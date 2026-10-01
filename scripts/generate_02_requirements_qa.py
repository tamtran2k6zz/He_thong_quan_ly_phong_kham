# -*- coding: utf-8 -*-
"""
Script sinh file 02_GenAI_SoftwareDevelopment_requirements-qa.docx
Thu thập và làm rõ yêu cầu nghiệp vụ thực tế giữa đội ngũ phát triển và khách hàng phòng khám (Nhóm 07).
Tuân thủ văn phong tự nhiên, câu hỏi nghiệp vụ đời thường, không dùng thuật ngữ kỹ thuật sâu.
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_style_helpers import (
    create_base_document,
    add_header_block,
    add_team_meta,
    add_h1,
    add_h2,
    add_p,
    add_code_block,
    format_row,
    set_table_borders,
    save_document,
    HEX_HEADER_BG,
    HEX_ROW_ALT
)

def build_requirements_qa():
    doc = create_base_document()

    add_header_block(
        doc,
        "THU THẬP, LÀM RÕ YÊU CẦU CỦA ỨNG DỤNG",
        "HỌC PHẦN: ỨNG DỤNG TRÍ TUỆ NHÂN TẠO"
    )

    add_team_meta(doc)

    add_p(
        doc,
        "Yêu cầu chức năng của một hệ thống là quan trọng vì yêu cầu cung cấp các cơ sở cho tất cả công việc phát triển hệ thống sau đó. "
        "Để xây dựng phần mềm phù hợp với thực tế hoạt động, nhóm phát triển đã tiến hành phỏng vấn trực tiếp ban quản lý, "
        "bác sĩ khám bệnh, nhân viên lễ tân và kế toán thu ngân tại phòng khám. "
        "Các câu hỏi tập trung vào quy trình tiếp đón người bệnh, việc khám bệnh, kê đơn thuốc, thanh toán viện phí và nhu cầu hỗ trợ công việc hàng ngày."
    )

    add_h1(doc, "1. Danh sách các câu hỏi khi thu thập và làm rõ yêu cầu của ứng dụng")
    add_p(
        doc,
        "Bảng tổng hợp 20 câu hỏi phỏng vấn khách hàng phòng khám và câu trả lời làm rõ yêu cầu thực tế:"
    )

    qa_natural_data = [
        (
            "1",
            "Hiện tại phòng khám tiếp nhận người bệnh đến khám bằng những hình thức nào?",
            "Người bệnh có thể đến đăng ký trực tiếp tại quầy lễ tân hoặc gọi điện thoại đặt hẹn trước. Phòng khám muốn phần mềm hỗ trợ quản lý cả hai hình thức này để không bị sót thông tin.",
            "Tiếp đón ban đầu"
        ),
        (
            "2",
            "Khi người bệnh đến khám lần đầu, nhân viên lễ tân cần thu thập những thông tin gì?",
            "Cần ghi lại họ tên, ngày sinh, giới tính, số căn cước công dân, số điện thoại, địa chỉ liên hệ và hỏi người bệnh có tiền sử dị ứng thuốc gì không để lưu vào hồ sơ.",
            "Hồ sơ người bệnh"
        ),
        (
            "3",
            "Phòng khám có tiếp nhận khám chữa bệnh bằng thẻ Bảo hiểm Y tế không?",
            "Có. Phòng khám tiếp nhận thẻ bảo hiểm y tế và cần kiểm tra mã thẻ để giảm trừ chi phí khám hoặc tiền thuốc cho người bệnh theo đúng quyền lợi được hưởng.",
            "Bảo hiểm Y tế"
        ),
        (
            "4",
            "Phòng khám có hay gặp tình trạng hai người bệnh bị xếp trùng một bác sĩ vào cùng một giờ không?",
            "Thỉnh thoảng có xảy ra do nhân viên ghi chép sổ tay. Phòng khám muốn phần mềm tự động chặn lại, không cho phép chọn giờ đã có người khác đặt khám với bác sĩ đó.",
            "Tránh trùng lịch"
        ),
        (
            "5",
            "Các bác sĩ tại phòng khám làm việc theo những khung giờ nào trong ngày?",
            "Phòng khám chia ba ca mỗi ngày gồm ca sáng từ 7h30 đến 11h30, ca chiều từ 13h30 đến 17h00 và ca tối từ 17h30 đến 20h30. Mỗi ngày sẽ có bác sĩ trực theo chuyên khoa riêng.",
            "Ca trực bác sĩ"
        ),
        (
            "6",
            "Khi người bệnh vào phòng khám, bác sĩ cần ghi chép những thông tin gì vào phiếu khám?",
            "Bác sĩ cần ghi triệu chứng bệnh nhân kể, đo các chỉ số cơ bản như huyết áp, nhịp tim, nhiệt độ, chiều cao, cân nặng, sau đó đưa ra kết luận chẩn đoán bệnh.",
            "Khám bệnh"
        ),
        (
            "7",
            "Phòng khám có cần ghi tên bệnh theo danh mục quy chuẩn nào không?",
            "Bác sĩ muốn chọn tên bệnh theo danh mục mã bệnh quốc tế ICD-10 của Bộ Y tế để phục vụ việc lưu trữ hồ sơ và kê đơn chuẩn xác.",
            "Chẩn đoán bệnh"
        ),
        (
            "8",
            "Bác sĩ kê đơn thuốc cho người bệnh như thế nào và lấy thuốc từ đâu?",
            "Bác sĩ chọn thuốc trực tiếp từ danh mục thuốc sẵn có tại nhà thuốc của phòng khám, ghi rõ số lượng, số lần uống trong ngày và hướng dẫn uống trước hay sau bữa ăn.",
            "Kê đơn thuốc"
        ),
        (
            "9",
            "Làm thế nào để phòng tránh việc bác sĩ kê nhầm loại thuốc mà người bệnh từng bị dị ứng?",
            "Phần mềm cần tự đối chiếu đơn thuốc vừa kê với thông tin dị ứng đã lưu trong hồ sơ bệnh nhân. Nếu có thuốc trùng hoạt chất dị ứng thì phải hiện cảnh báo ngay trên màn hình.",
            "An toàn dùng thuốc"
        ),
        (
            "10",
            "Sau khi khám xong, quy trình thanh toán tiền khám và tiền thuốc diễn ra tại đâu?",
            "Người bệnh cầm phiếu khám ra quầy kế toán thu ngân. Kế toán sẽ xem tổng chi phí, trừ tiền bảo hiểm chi trả nếu có, rồi thu phần tiền còn lại của người bệnh.",
            "Thanh toán viện phí"
        ),
        (
            "11",
            "Người bệnh có thể thanh toán bằng những cách nào ngoài tiền mặt?",
            "Nhiều người bệnh muốn chuyển khoản ngân hàng. Phòng khám muốn phần mềm tạo sẵn mã QR có sẵn số tiền để người bệnh quét mã trả tiền cho nhanh, tránh nhầm lẫn số tiền.",
            "Hình thức trả tiền"
        ),
        (
            "12",
            "Người bệnh có cần nhận giấy tờ gì sau khi đã nộp tiền không?",
            "Kế toán cần in phiếu thu tiền viện phí có đầy đủ danh mục tiền khám, tiền thuốc và số tiền thực thu để người bệnh giữ làm biên lai.",
            "Phiếu thu tiền"
        ),
        (
            "13",
            "Phòng khám có những bộ phận nhân viên nào sử dụng phần mềm hàng ngày?",
            "Có bốn bộ phận gồm nhân viên lễ tân tiếp đón, các bác sĩ khám bệnh, nhân viên kế toán thu ngân và người quản lý phòng khám.",
            "Người dùng phần mềm"
        ),
        (
            "14",
            "Nhân viên lễ tân hoặc kế toán có được xem chi tiết kết quả khám và bệnh án của bệnh nhân không?",
            "Không. Lễ tân chỉ cần biết thông tin cá nhân và giờ hẹn. Kế toán chỉ cần xem bảng tính tiền. Chi tiết triệu chứng và kết quả bệnh án chỉ dành riêng cho bác sĩ để bảo đảm tính riêng tư.",
            "Bảo mật bệnh án"
        ),
        (
            "15",
            "Bác sĩ có thường xuyên cần xem lại các lần khám trước của người bệnh không?",
            "Rất cần. Bệnh nhân mạn tính như tiểu đường, huyết áp thường tái khám nhiều lần. Bác sĩ cần xem nhanh lần trước người bệnh uống thuốc gì, chỉ số huyết áp ra sao để điều chỉnh liều lượng.",
            "Lịch sử khám cũ"
        ),
        (
            "16",
            "Phòng khám mong muốn trợ lý ảo hỗ trợ việc gì cho bác sĩ trong lúc khám?",
            "Mong muốn trợ lý ảo tự đọc các lần khám cũ rồi tóm tắt ngắn gọn các bệnh nền, loại thuốc đang uống để bác sĩ liếc mắt là nắm được ngay trong mười giây đầu tiên.",
            "Hỗ trợ tóm tắt"
        ),
        (
            "17",
            "Phòng khám có cho phép trợ lý ảo tự đưa ra kết luận chẩn đoán hay kê đơn thuốc thay bác sĩ không?",
            "Tuyệt đối không. Trí tuệ nhân tạo chỉ sắp xếp và tóm tắt thông tin hỗ trợ. Quyết định chẩn đoán và kê đơn bắt buộc phải do bác sĩ trực tiếp chịu trách nhiệm.",
            "Trách nhiệm y tế"
        ),
        (
            "18",
            "Người bệnh sau khi khám xong thường gặp khó khăn gì khi dùng thuốc tại nhà?",
            "Nhiều người hay quên giờ uống thuốc hoặc uống sai bữa. Bác sĩ mong muốn phần mềm tự tạo một bản hướng dẫn dặn dò dễ hiểu, chia rõ sáng uống viên nào, trưa uống viên nào và hẹn ngày tái khám.",
            "Dặn dò dùng thuốc"
        ),
        (
            "19",
            "Người bệnh thường hỏi những câu hỏi gì khi liên hệ với phòng khám trước khi đến khám?",
            "Họ hay hỏi về giờ mở cửa, chi phí khám các khoa, thủ tục mang giấy tờ bảo hiểm ra sao. Nếu có một hộp trả lời tự động giải đáp các câu này trên máy thì lễ tân sẽ đỡ mất thời gian trực điện thoại.",
            "Hỏi đáp thường gặp"
        ),
        (
            "20",
            "Người quản lý phòng khám cần xem những số liệu tổng kết nào vào cuối ngày hoặc cuối tháng?",
            "Cần xem tổng số lượt người bệnh đến khám theo từng khoa, doanh thu tiền khám, doanh thu tiền thuốc và số lượng thuốc còn tồn trong kho để kịp thời nhập thêm.",
            "Báo cáo quản lý"
        )
    ]

    table = doc.add_table(rows=len(qa_natural_data) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(0.6), Inches(2.2), Inches(2.7), Inches(1.0)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    headers = ["STT", "Câu hỏi trao đổi với khách hàng", "Ý kiến phản hồi của phòng khám", "Nội dung"]
    for i, h_text in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h_text
    format_row(table.rows[0], HEX_HEADER_BG, RGBColor(255, 255, 255), bold=True, is_header=True)

    for r_idx, data in enumerate(qa_natural_data):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(data):
            row.cells[c_idx].text = val
        bg = HEX_ROW_ALT if (r_idx % 2 == 1) else "FFFFFF"
        format_row(row, bg, RGBColor(15, 23, 42), bold=False, is_header=False)

    set_table_borders(table)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    add_h1(doc, "2. Yêu cầu chức năng và phi chức năng của ứng dụng")

    add_h2(doc, "2.1. Yêu cầu chức năng theo nhu cầu phòng khám")
    add_p(
        doc,
        "Dựa trên ý kiến trao đổi với khách hàng, các yêu cầu phục vụ công việc hàng ngày của phòng khám gồm có:"
    )

    fr_natural = [
        ("Nghiệp vụ Tiếp đón và Lịch hẹn: ", "Ghi nhận thông tin người bệnh, tìm kiếm nhanh theo số điện thoại hoặc căn cước công dân, xếp lịch khám cho người bệnh và tự động ngăn chặn việc đặt trùng giờ của bác sĩ."),
        ("Nghiệp vụ Khám bệnh và Hồ sơ sức khỏe: ", "Mở sổ khám bệnh trên máy tính, ghi nhận huyết áp, cân nặng, nhịp tim, tự động tính thể trạng người bệnh, chọn tên bệnh và lưu trữ lịch sử các lần khám."),
        ("Nghiệp vụ Kê đơn và Nhà thuốc: ", "Chọn thuốc có sẵn trong tủ thuốc phòng khám, kiểm tra và cảnh báo nếu người bệnh bị dị ứng với thuốc định kê, in đơn thuốc rõ ràng cho bệnh nhân."),
        ("Nghiệp vụ Viện phí và Thu ngân: ", "Tính tổng tiền khám và tiền thuốc, tự động trừ phần tiền bảo hiểm y tế chi trả, hiển thị mã quét chuyển khoản ngân hàng và in biên lai nộp tiền."),
        ("Trợ lý ảo hỗ trợ công việc: ", "Tự động đọc tóm tắt các lần khám cũ cho bác sĩ xem nhanh, giải đáp các thắc mắc thông thường của người bệnh về giấy tờ và tự động in bảng dặn dò giờ uống thuốc."),
        ("Quản lý và Báo cáo: ", "Phân chia ca trực cho bác sĩ, kiểm tra số lượng thuốc còn tồn trong kho và xem báo cáo doanh thu phòng khám theo ngày hoặc theo tháng.")
    ]
    for title, desc in fr_natural:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r2 = p.add_run(desc)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)

    add_h2(doc, "2.2. Yêu cầu phi chức năng phục vụ trải nghiệm người dùng")
    nfr_natural = [
        ("Dễ sử dụng và thuận tiện: ", "Giao diện rõ ràng, chữ viết và số liệu to rõ, nhân viên lễ tân và bác sĩ lớn tuổi đều có thể thao tác nhanh mà không cần học nhiều."),
        ("Tốc độ xử lý nhanh: ", "Tìm kiếm thông tin bệnh nhân và mở hồ sơ diễn ra ngay lập tức, không để người bệnh phải chờ đợi lâu tại quầy tiếp đón."),
        ("Bảo mật và Riêng tư: ", "Thông tin cá nhân và bệnh tật của người bệnh được bảo vệ an toàn, nhân viên bộ phận nào chỉ được xem thông tin thuộc nhiệm vụ của bộ phận đó."),
        ("Hoạt động ổn định: ", "Phần mềm chạy bền bỉ trên máy tính phòng khám, có bản lưu trữ dự phòng dữ liệu để tránh mất mát khi xảy ra sự cố mất điện."),
        ("Ranh giới y tế an toàn: ", "Trợ lý ảo chỉ đóng vai trò hỗ trợ dọn dẹp giấy tờ và nhắc nhở, mọi chỉ định dùng thuốc và kết luận bệnh tật đều do bác sĩ quyết định.")
    ]
    for title, desc in nfr_natural:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r2 = p.add_run(desc)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)

    add_h1(doc, "3. Sơ đồ phân cấp chức năng của ứng dụng (Functional Decomposition Diagram)")
    add_p(
        doc,
        "Sơ đồ phân cấp chức năng thể hiện trực quan các mảng công việc của phòng khám được đưa vào phần mềm:"
    )

    fdd_natural_ascii = """+----------------------------------------------------------------------------------------------------+
|                         PHẦN MỀM QUẢN LÝ PHÒNG KHÁM ĐA KHOA CÓ TRỢ LÝ ẢO                            |
+----------------------------------------------------------------------------------------------------+
       |                                |                               |                           |
       v                                v                               v                           v
+------------------+         +--------------------+         +--------------------+       +--------------------+
| 1. TIẾP ĐÓN VÀ   |         | 2. KHÁM BỆNH VÀ    |         | 3. NHÀ THUỐC VÀ    |       | 4. TRỢ LÝ ẢO VÀ    |
|    ĐẶT LỊCH HẸN  |         |    HỒ SƠ BỆNH ÁN   |         |    THU VIỆN PHÍ    |       |    BÁO CÁO         |
+------------------+         +--------------------+         +--------------------+       +--------------------+
       |                                |                               |                           |
       +-- 1.1 Tiếp nhận bệnh nhân,     +-- 2.1 Hàng đợi chờ khám       +-- 3.1 Kê đơn thuốc từ     +-- 4.1 Tóm tắt lịch sử
       |   tìm kiếm theo SĐT, CCCD      |   theo số thứ tự              |   danh mục kho thuốc      |   khám cũ cho bác sĩ
       |                                |                               |                           |
       +-- 1.2 Đăng ký hồ sơ mới        +-- 2.2 Đo chỉ số sinh hiệu     +-- 3.2 Cảnh báo dị ứng     +-- 4.2 Hộp giải đáp
       |   và thông tin bảo hiểm        |   (huyết áp, tim, cân nặng)   |   khi kê thuốc            |   thắc mắc người bệnh
       |                                |                               |                           |
       +-- 1.3 Đặt lịch khám và         +-- 2.3 Tự động tính thể trạng  +-- 3.3 Tính tiền viện phí  +-- 4.3 Tạo giấy dặn dò
       |   chọn bác sĩ chuyên khoa      |   người bệnh (BMI)            |   và trừ bảo hiểm y tế    |   giờ uống thuốc
       |                                |                               |                           |
       +-- 1.4 Tự động chặn trùng       +-- 2.4 Chẩn đoán bệnh theo     +-- 3.4 Tạo mã quét ngân    +-- 4.4 Báo cáo doanh thu
       |   giờ khám của bác sĩ          |   danh mục chuẩn Bộ Y tế      |   hàng chuyển khoản       |   và số lượt khám
       |                                |                               |                           |
       +-- 1.5 Xếp lịch trực bác sĩ     +-- 2.5 Lưu hồ sơ điện tử       +-- 3.5 In phiếu thu tiền   +-- 4.5 Kiểm tra lượng
           theo ca trong ngày               và chuyển sang thu phí          cho người bệnh              thuốc còn trong kho
+----------------------------------------------------------------------------------------------------+"""

    add_code_block(doc, fdd_natural_ascii)

    save_document(doc, "02_GenAI_SoftwareDevelopment_requirements-qa.docx")

if __name__ == "__main__":
    build_requirements_qa()
