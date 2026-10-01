# -*- coding: utf-8 -*-
"""
Bộ công cụ định dạng và hỗ trợ sinh tài liệu Word chuẩn cho 7 giai đoạn SDLC.
Tuân thủ quy chuẩn văn phong chuyên nghiệp:
- Không dùng dấu gạch ngang dài (dùng dấu phẩy hoặc câu đơn).
- Không dùng từ nối sáo rỗng, cấu trúc rập khuôn.
- Dùng động từ trực tiếp và dữ liệu kỹ thuật cụ thể.
"""

import os
import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate 900
COLOR_TEAL = RGBColor(13, 148, 136)       # Teal 600
COLOR_MUTED = RGBColor(71, 85, 105)       # Slate 600
COLOR_DARK_BLUE = RGBColor(30, 58, 138)   # Blue 900

HEX_HEADER_BG = "1E293B"                 # Slate 800
HEX_ROW_ALT = "F8FAFC"                   # Slate 50
HEX_CODE_BG = "F1F5F9"                   # Slate 100
HEX_BORDER = "CBD5E1"                    # Slate 300

def create_base_document():
    """Khởi tạo tài liệu với lề 1.0 inch và thiết lập trang chuẩn A4."""
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
    return doc

def set_cell_background(cell, fill_hex):
    """Đặt màu nền cho ô trong bảng."""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Đặt khoảng đệm trong ô của bảng."""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def set_table_borders(table, color="CBD5E1"):
    """Thiết lập viền mỏng thanh lịch cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def format_row(row, bg_hex, text_color, bold=False, is_header=False):
    """Định dạng văn bản và màu nền cho một hàng trong bảng."""
    for cell in row.cells:
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = text_color
                run.font.bold = bold
                run.font.size = Pt(10) if is_header else Pt(9.5)

def add_header_block(doc, title, subtitle=None):
    """Tạo phần tiêu đề đầu trang theo mẫu môn học."""
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO, TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_MUTED

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(8)
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.name = "Times New Roman"
        r_sub.font.size = Pt(12)
        r_sub.font.bold = True
        r_sub.font.color.rgb = COLOR_TEAL

def add_team_meta(doc, show_members=True):
    """Ghi nhận thông tin nhóm sinh viên thực hiện đề tài."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2

    runs_data = [
        ("Đơn vị thực hiện: ", True), ("Nhóm 03 (Lớp An toàn thông tin K23A, ICTU)\n", False),
        ("Thành viên nhóm:\n", True),
        ("1. Đinh Gia Bảo ", False), ("(Trưởng nhóm, Kiến trúc hệ thống và AI Backend)\n", True),
        ("2. Trần Đặng Công Tâm ", False), ("(Thành viên chính, Giao diện người dùng và Kiểm thử QA)\n", True),
        ("Tên ứng dụng: ", True), ("Hệ thống quản lý phòng khám có tích hợp AI (CMS-AI)\n", False),
        ("Thời gian thực hiện: ", True), ("Từ 27/07/2026 đến 27/09/2026 (Thời lượng 9 tuần)", False)
    ]
    for text, bold in runs_data:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = bold
        r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_h1(doc, text):
    """Tiêu đề cấp 1."""
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return h

def add_h2(doc, text):
    """Tiêu đề cấp 2."""
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return h

def add_h3(doc, text):
    """Tiêu đề cấp 3."""
    h = doc.add_heading(level=3)
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(2)
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_MUTED
    return h

def add_p(doc, text):
    """Đoạn văn thông thường với giãn dòng chuẩn."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_PRIMARY
    return p

def add_code_block(doc, text):
    """Hộp văn bản mã nguồn hoặc sơ đồ cấu trúc."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def save_document(doc, filename):
    """Lưu file đồng thời vào thư mục Desktop của người dùng và thư mục lưu trữ dự án."""
    desktop_dir = r"C:\Users\MSI\OneDrive\Desktop\CacGiaiDoanThucHien"
    project_dir = r"d:\ICTU\Nam 3\ICTU_2026-2027\Ứng dụng trí tuệ nhân tạo - Project\He_thong_quan_ly_phong_kham\docs\7_Giai_Doan_Hoan_Thien"
    
    os.makedirs(desktop_dir, exist_ok=True)
    os.makedirs(project_dir, exist_ok=True)

    dest_desktop = os.path.join(desktop_dir, filename)
    dest_project = os.path.join(project_dir, filename)

    try:
        doc.save(dest_desktop)
        print(f"-> Đã lưu thành công vào Desktop: {filename}")
    except PermissionError:
        print(f"! Cảnh báo: Tệp {filename} đang mở trong Word trên Desktop, hãy đóng tệp để cập nhật bản mới.")

    try:
        doc.save(dest_project)
        print(f"-> Đã lưu thành công vào kho lưu trữ: {filename}")
    except PermissionError:
        print(f"! Cảnh báo: Tệp {filename} đang mở trong Word trong dự án.")
