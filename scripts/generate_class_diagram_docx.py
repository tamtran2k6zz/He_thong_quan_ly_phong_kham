# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) Tài Liệu Thiết Kế Hướng Đối Tượng (Mô Hình Lớp)
cho Nhóm 07 - Hệ thống quản lý phòng khám có tích hợp AI.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    """Đặt màu nền cho ô trong bảng"""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Đặt lề trong cho ô của bảng"""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def format_row(row, bg_hex, text_color, bold=False, is_header=False):
    """Định dạng cho toàn bộ hàng trong bảng"""
    for cell in row.cells:
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = text_color
                run.font.bold = bold
                if is_header:
                    run.font.size = Pt(10)
                    run.font.bold = True
                else:
                    run.font.size = Pt(9.5)

def add_code_block(doc, text):
    """Thêm một khối code định dạng với nền xám nhạt"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")  # Slate 100
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_h1(doc, text):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(13, 148, 136) # Teal 600
    return h

def add_h2(doc, text):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    return h

def add_h3(doc, text):
    h = doc.add_heading(level=3)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    return h

def add_method_spec(doc, name, desc, inputs, output, flow, start_cond, end_cond):
    """Đặc tả một phương thức theo đúng mẫu yêu cầu"""
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(4)
    p_name.paragraph_format.space_after = Pt(2)
    r_bullet = p_name.add_run("o Tên phương thức: ")
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = RGBColor(13, 148, 136)
    r_val = p_name.add_run(name)
    r_val.font.bold = True
    r_val.font.name = "Consolas"
    
    # Mô tả
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("   • Mô tả: ").font.bold = True
    p.add_run(desc)
    
    # Tham số đầu vào
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("   • Tham số đầu vào: ").font.bold = True
    p.add_run(inputs)
    
    # Kết quả đầu ra
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("   • Kết quả đầu ra: ").font.bold = True
    p.add_run(output)
    
    # Luồng xử lý
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("   • Luồng xử lý: ").font.bold = True
    p.add_run(flow)
    
    # Điều kiện bắt đầu
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("   • Điều kiện bắt đầu: ").font.bold = True
    p.add_run(start_cond)
    
    # Điều kiện kết thúc
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("   • Điều kiện kết thúc: ").font.bold = True
    p.add_run(end_cond)

def build_docx():
    doc = docx.Document()

    # Thiết lập lề 1.0 inch (2.54 cm)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        # Header & Footer
        footer = s.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("TÀI LIỆU THIẾT KẾ HƯỚNG ĐỐI TƯỢNG — NHÓM 07 — ĐINH GIA BẢO & TRẦN ĐẶNG CÔNG TÂM")
        r_ft.font.name = "Times New Roman"
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # TRANG BÌA VÀ THÔNG TIN ĐỒ ÁN
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(71, 85, 105)
    p_inst.paragraph_format.space_after = Pt(24)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("TÀI LIỆU THIẾT KẾ HƯỚNG ĐỐI TƯỢNG (MÔ HÌNH LỚP)\n")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(16)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(13, 148, 136) # Teal 600

    r_t2 = p_title.add_run("HỆ THỐNG QUẢN LÝ PHÒNG KHÁM CÓ TÍCH HỢP AI")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(17)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_after = Pt(20)

    # Bảng thông tin nhóm và thời gian thực hiện
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    info_table.columns[0].width = Inches(2.2)
    info_table.columns[1].width = Inches(4.3)
    
    rows_data = [
        ("Nhóm thực hiện:", "Nhóm 07"),
        ("Thành viên nhóm:", "1. Đinh Gia Bảo (Trưởng nhóm — Kiến trúc & AI Lead)\n2. Trần Đặng Công Tâm (Thành viên — Fullstack Dev & QA Lead)"),
        ("Tên ứng dụng:", "Hệ thống quản lý phòng khám có tích hợp AI (CMS-AI)"),
        ("Thời gian thực hiện:", "Từ 27/07/2026 đến 27/09/2026 (9 tuần)")
    ]
    for idx, (label, val) in enumerate(rows_data):
        row = info_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, 80, 80, 120, 120)
        set_cell_margins(cell_val, 80, 80, 120, 120)
        
        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.name = "Times New Roman"
        r1.font.bold = True
        r1.font.size = Pt(10)
        
        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # =========================================================================
    # MỤC 1: MÔ HÌNH LỚP (CLASS DIAGRAM)
    # =========================================================================
    add_h1(doc, "1. MÔ HÌNH LỚP (CLASS DIAGRAM)")
    p = doc.add_paragraph()
    p.add_run("Mô hình lớp (Class Diagram) tổng thể của ")
    p.add_run("Hệ thống quản lý phòng khám có tích hợp AI").font.bold = True
    p.add_run(" mô tả chi tiết cấu trúc tĩnh của hệ thống, bao gồm các thực thể dữ liệu nghiệp vụ (Domain Entities), các dịch vụ điều phối thuật toán (Core Services), và hệ sinh thái Trợ lý AI Hành chính 3 lớp:")

    add_code_block(doc,
"""+-----------------------------------------------------------------------------------------------------------------------+
|                                        MÔ HÌNH LỚP TỔNG THỂ HỆ THỐNG PHÒNG KHÁM                                       |
+-----------------------------------------------------------------------------------------------------------------------+

  [User] 1 ────── 0..1 [Doctor] 1 ────── 0..* [Shift]
    │                     │
    │ 1                   │ 1
    │                     ▼
    │ 0..*             [Appointment] 0..* ────── 1 [Clinic] 0..* ────── 1 [Specialty]
    │                     ▲                                                    ▲
    │                     │ 0..*                                               │ 1
    ▼                     │                                                    │
  [AuditLog]           [Patient] 1 ──────────────────────── 0..* [Doctor] ─────┘
                          │ 1
                          ├─────────────────┐
                          ▼ 0..*            ▼ 0..*
                   [MedicalRecord]       [Invoice]
                    1 │       │ 1           ▲
                      │       │             │ 0..1
             0..* ┌───┘       └───┐ 0..1    │
                  ▼               ▼         │
            [ServiceOrder]   [Prescription] ┘
                                  │ 1
                                  ▼ 1..*
                           [PrescriptionItem] 0..* ────── 1 [Medicine]

  +-----------------------+     +-----------------------+     +-----------------------+
  |    ConflictChecker    |     |     PIIAnonymizer     |     |   AdminAIGuardrails   |
  |  (Core Math Service)  |     |  (Privacy Protection) |     |  (Medical Safety L2)  |
  +-----------------------+     +-----------------------+     +-----------------------+
              │                             │                             │
              ▼ validates                   ▼ anonymizes                  ▼ filters
        [Appointment]              [AdminAIService] <─────────────────────┘
                                   (Gemini Live & Offline Mock)
+-----------------------------------------------------------------------------------------------------------------------+""")

    p_note = doc.add_paragraph()
    p_note.add_run("Biểu đồ lớp UML chi tiết biểu diễn bằng chuẩn ngôn ngữ mô hình hóa Mermaid:")
    p_note.runs[0].font.italic = True

    add_code_block(doc,
"""classDiagram
    direction TB

    class User {
        +int id
        +string username
        +string email
        +string full_name
        +string hashed_password
        +string role
        +bool is_active
        +datetime created_at
        +verify_password(string plain_password) bool
        +has_permission(string required_role) bool
        +deactivate() void
    }

    class Patient {
        +int id
        +string medical_code
        +string full_name
        +date date_of_birth
        +string gender
        +string phone
        +string identity_card
        +string address
        +string insurance_number
        +string medical_history
        +string drug_allergies
        +string emergency_contact
        +datetime created_at
        +calculate_age() int
        +check_drug_allergy(string medicine_name) bool
        +update_contact(string new_phone, string new_address) void
    }

    class Doctor {
        +int id
        +int user_id
        +int specialty_id
        +int clinic_id
        +string title
        +string bio
        +string phone
        +get_schedule(date query_date) list
        +is_available(datetime start_time, datetime end_time) bool
    }

    class Appointment {
        +int id
        +string appointment_code
        +int patient_id
        +int doctor_id
        +int clinic_id
        +date appointment_date
        +time start_time
        +time end_time
        +string status
        +string reason
        +string notes
        +datetime created_at
        +confirm_booking() void
        +check_in() void
        +cancel_appointment(string cancel_reason) void
        +reschedule(date new_date, time new_start, time new_end) bool
    }

    class MedicalRecord {
        +int id
        +string record_code
        +int patient_id
        +int doctor_id
        +int appointment_id
        +datetime exam_date
        +string chief_complaint
        +string blood_pressure
        +int heart_rate
        +float temperature
        +int respiratory_rate
        +float weight
        +float height
        +float bmi
        +string physical_exam
        +string diagnosis_icd10
        +string icd10_code
        +string doctor_notes
        +string status
        +datetime created_at
        +calculate_bmi() float
        +add_service_order(string service_code, string service_name, float price) ServiceOrder
        +create_prescription(string advice) Prescription
        +complete_examination() void
    }

    class Medicine {
        +int id
        +string code
        +string name
        +string active_ingredient
        +string dosage_form
        +string unit
        +float unit_price
        +int stock_quantity
        +string usage_instructions
        +bool is_active
        +deduct_stock(int qty) bool
        +restock(int qty) void
        +is_in_stock(int required_qty) bool
    }

    class Prescription {
        +int id
        +string prescription_code
        +int medical_record_id
        +int doctor_id
        +int patient_id
        +string diagnosis
        +string advice
        +datetime created_at
        +add_item(int medicine_id, int quantity, string dosage, string frequency, int duration_days, string instructions) PrescriptionItem
        +calculate_total_medicine_cost() float
    }

    class Invoice {
        +int id
        +string invoice_code
        +int medical_record_id
        +int patient_id
        +float consultation_fee
        +float service_fee
        +float medicine_fee
        +float total_amount
        +float insurance_discount
        +float patient_pay_amount
        +string payment_status
        +string payment_method
        +string transaction_code
        +int cashier_id
        +datetime paid_at
        +string notes
        +datetime created_at
        +calculate_total_bill(float bhyt_rate) void
        +generate_vietqr_payload() dict
        +process_payment(string method, string tx_code, int cashier_user_id) bool
    }

    class ConflictChecker {
        +check_appointment_conflict(int doctor_id, int clinic_id, date apt_date, time start_time, time end_time, int exclude_id) tuple
        +check_shift_alignment(int doctor_id, date apt_date, time start_time, time end_time) bool
    }

    class AdminAIService {
        -PIIAnonymizer anonymizer
        -AdminAIGuardrails guardrails
        +generate_pre_visit_summary(int patient_id) dict
        +answer_clinic_faq(string question) dict
        +generate_discharge_instructions(int medical_record_id) dict
    }

    User "1" *-- "0..1" Doctor : doctor_profile
    Patient "1" --> "0..*" Appointment : books
    Doctor "1" --> "0..*" Appointment : attends
    Appointment "1" --> "0..1" MedicalRecord : leads_to
    MedicalRecord "1" *-- "0..1" Prescription : prescribes
    MedicalRecord "1" --> "0..1" Invoice : generates
    ConflictChecker ..> Appointment : validates
    AdminAIService ..> MedicalRecord : summarizes""")

    # =========================================================================
    # MỤC 2: ĐẶC TẢ CHI TIẾT CLASS
    # =========================================================================
    add_h1(doc, "2. ĐẶC TẢ CLASS (CLASS SPECIFICATION)")
    
    # -------------------------------------------------------------------------
    # 2.1 USER
    # -------------------------------------------------------------------------
    add_h2(doc, "2.1. Lớp User (Người dùng & Xác thực Phân quyền RBAC)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Đại diện cho tài khoản nhân sự phòng khám (Admin, Receptionist, Doctor, Accountant), thực hiện định danh JWT, bảo mật mật khẩu băm bcrypt và kiểm soát phân quyền RBAC.")

    add_h3(doc, "Các thuộc tính của lớp User:")
    user_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("id", "Integer", "4 bytes, PK, Auto Increment", "Định danh duy nhất của người dùng"),
        ("username", "String(50)", "Tối đa 50 ký tự, UNIQUE, NOT NULL, Index", "Tên đăng nhập hệ thống"),
        ("email", "String(100)", "Tối đa 100 ký tự, UNIQUE, Index, NULLABLE", "Địa chỉ email liên hệ"),
        ("full_name", "String(100)", "Tối đa 100 ký tự, NOT NULL", "Họ và tên đầy đủ của nhân sự"),
        ("hashed_password", "String(255)", "Tối đa 255 ký tự, NOT NULL", "Mật khẩu băm theo chuẩn bcrypt"),
        ("role", "String(20)", "Tối đa 20 ký tự, NOT NULL", "Vai trò: admin, receptionist, doctor, accountant"),
        ("is_active", "Boolean", "1 byte, DEFAULT True, NOT NULL", "Trạng thái hoạt động của tài khoản"),
        ("created_at", "DateTime", "8 bytes, DEFAULT utcnow, NOT NULL", "Thời điểm khởi tạo tài khoản")
    ]
    tbl = doc.add_table(rows=len(user_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.2)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(2.0)
    for idx, (c1, c2, c3, c4) in enumerate(user_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp User:")
    
    add_method_spec(
        doc,
        name="verify_password",
        desc="Kiểm tra tính chính xác của mật khẩu văn bản thô do người dùng nhập vào so với chuỗi băm bcrypt đã lưu trong CSDL.",
        inputs="plain_password: String, kích thước tối đa 128 ký tự (Mật khẩu người dùng nhập từ form đăng nhập).",
        output="is_valid: Boolean, kích thước 1 byte (True nếu mật khẩu trùng khớp, False nếu sai mật khẩu).",
        flow="1. Trích xuất chuỗi băm self.hashed_password.\n2. Sử dụng thư viện passlib giải thuật bcrypt để so khớp chuỗi thô plain_password với chuỗi băm.\n3. Trả về kết quả True hoặc False.",
        start_cond="Bản ghi User đã được nạp từ CSDL và self.hashed_password có giá trị hợp lệ.",
        end_cond="Trả về kết quả xác thực mật khẩu mà không làm thay đổi trạng thái đối tượng."
    )

    add_method_spec(
        doc,
        name="has_permission",
        desc="Kiểm tra người dùng hiện tại có đủ thẩm quyền để truy cập tính năng yêu cầu vai trò cụ thể hay không.",
        inputs="required_role: String, kích thước tối đa 20 ký tự (Vai trò tối thiểu được yêu cầu).",
        output="has_access: Boolean, kích thước 1 byte (True nếu được phép truy cập, False nếu bị từ chối).",
        flow="1. Kiểm tra tài khoản có đang hoạt động không (self.is_active == True). Nếu không, trả về False.\n2. Nếu self.role == 'admin' thì luôn trả về True (Admin có toàn quyền quản trị).\n3. Nếu self.role == required_role thì trả về True.\n4. Ngược lại, trả về False.",
        start_cond="Đối tượng User đã được xác thực danh tính qua JWT Token.",
        end_cond="Xác định quyền truy cập của người dùng tại Endpoint tương ứng."
    )

    add_method_spec(
        doc,
        name="deactivate",
        desc="Khóa tài khoản nhân viên khi nghỉ việc hoặc tạm đình chỉ, ngăn chặn đăng nhập hệ thống.",
        inputs="Không có.",
        output="None, kích thước 0 byte.",
        flow="1. Gán giá trị self.is_active = False.",
        start_cond="Tài khoản đang ở trạng thái is_active = True.",
        end_cond="Thuộc tính is_active chuyển thành False, chặn các phiên làm việc tiếp theo."
    )

    # -------------------------------------------------------------------------
    # 2.2 PATIENT
    # -------------------------------------------------------------------------
    add_h2(doc, "2.2. Lớp Patient (Hồ sơ Bệnh nhân)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Quản lý toàn bộ thông tin nhân khẩu học, mã số định danh y tế, thẻ BHYT, tiền sử bệnh nền và danh sách dị ứng thuốc của người bệnh.")

    add_h3(doc, "Các thuộc tính của lớp Patient:")
    patient_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("id", "Integer", "4 bytes, PK, Auto Increment", "Định danh duy nhất bệnh nhân"),
        ("medical_code", "String(30)", "Tối đa 30 ký tự, UNIQUE, NOT NULL, Index", "Mã bệnh nhân: BN-YYYYMMDD-XXXX"),
        ("full_name", "String(100)", "Tối đa 100 ký tự, NOT NULL, Index", "Họ và tên người bệnh"),
        ("date_of_birth", "Date", "4 bytes, NOT NULL", "Ngày tháng năm sinh"),
        ("gender", "String(10)", "Tối đa 10 ký tự, NOT NULL", "Giới tính: Nam, Nữ, Khác"),
        ("phone", "String(20)", "Tối đa 20 ký tự, NOT NULL, Index", "Số điện thoại liên lạc"),
        ("identity_card", "String(20)", "Tối đa 20 ký tự, UNIQUE, Index, NULLABLE", "Số thẻ CCCD / CMND 12 chữ số"),
        ("address", "String(255)", "Tối đa 255 ký tự, NULLABLE", "Địa chỉ nơi ở hiện tại"),
        ("insurance_number", "String(25)", "Tối đa 25 ký tự, Index, NULLABLE", "Số thẻ Bảo hiểm Y tế (15 ký tự)"),
        ("medical_history", "Text", "Tối đa 65,535 ký tự, NULLABLE", "Tiền sử bệnh nền mạn tính"),
        ("drug_allergies", "Text", "Tối đa 65,535 ký tự, NULLABLE", "Tiền sử dị ứng thuốc và thực phẩm"),
        ("emergency_contact", "String(150)", "Tối đa 150 ký tự, NULLABLE", "Thông tin người thân báo tin khẩn cấp"),
        ("created_at", "DateTime", "8 bytes, DEFAULT utcnow, NOT NULL", "Ngày giờ tạo hồ sơ bệnh án")
    ]
    tbl = doc.add_table(rows=len(patient_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(1.9)
    for idx, (c1, c2, c3, c4) in enumerate(patient_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp Patient:")
    
    add_method_spec(
        doc,
        name="calculate_age",
        desc="Tính toán số tuổi hiện tại của bệnh nhân tính theo năm dựa trên ngày tháng năm sinh.",
        inputs="Không có.",
        output="age: Integer, kích thước 4 bytes (Số tuổi nguyên dương).",
        flow="1. Lấy ngày hiện tại: today = date.today().\n2. Tính độ lệch năm: age = today.year - self.date_of_birth.year.\n3. Nếu chưa tới ngày sinh nhật trong năm hiện tại: if (today.month, today.day) < (self.date_of_birth.month, self.date_of_birth.day): age -= 1.\n4. Trả về age.",
        start_cond="Thuộc tính date_of_birth có giá trị ngày hợp lệ.",
        end_cond="Trả về số tuổi chính xác của người bệnh."
    )

    add_method_spec(
        doc,
        name="check_drug_allergy",
        desc="Kiểm tra tên thuốc bác sĩ định kê đơn có nằm trong tiền sử dị ứng đã ghi nhận của bệnh nhân hay không.",
        inputs="medicine_name: String, kích thước tối đa 150 ký tự (Tên thương mại hoặc hoạt chất thuốc cần kiểm tra).",
        output="has_allergy: Boolean, kích thước 1 byte (True nếu có nguy cơ dị ứng, False nếu an toàn).",
        flow="1. Nếu self.drug_allergies rỗng hoặc None, trả về False.\n2. Chuyển medicine_name và self.drug_allergies về chữ thường (lowercase).\n3. Kiểm tra xem chuỗi con có xuất hiện: medicine_name.lower() in self.drug_allergies.lower().\n4. Trả về kết quả kiểm tra.",
        start_cond="Có tên thuốc cần đối soát.",
        end_cond="Cảnh báo nguy cơ sốc phản vệ hoặc phản ứng dị ứng thuốc trước khi kê đơn."
    )

    add_method_spec(
        doc,
        name="update_contact",
        desc="Cập nhật số điện thoại và địa chỉ cư trú mới của người bệnh khi có thay đổi.",
        inputs="new_phone: String, kích thước tối đa 20 ký tự (Số điện thoại mới).\nnew_address: String, kích thước tối đa 255 ký tự (Địa chỉ nơi ở mới).",
        output="None, kích thước 0 byte.",
        flow="1. Kiểm tra tính hợp lệ của số điện thoại new_phone.\n2. Cập nhật self.phone = new_phone.\n3. Cập nhật self.address = new_address.",
        start_cond="Nhận được thông tin liên hệ mới từ tiếp đón hoặc bệnh nhân.",
        end_cond="Thuộc tính phone và address của đối tượng được cập nhật thành công."
    )

    # -------------------------------------------------------------------------
    # 2.3 DOCTOR
    # -------------------------------------------------------------------------
    add_h2(doc, "2.3. Lớp Doctor (Bác sĩ Chuyên khoa)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Quản lý bác sĩ khám bệnh, liên kết 1-1 với tài khoản User và tham chiếu tới buồng khám bệnh, chuyên khoa và ca trực.")

    add_h3(doc, "Các thuộc tính của lớp Doctor:")
    doc_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("id", "Integer", "4 bytes, PK, Auto Increment", "Định danh duy nhất bác sĩ"),
        ("user_id", "Integer", "4 bytes, FK (users.id), UNIQUE, NOT NULL", "Tài khoản người dùng liên kết"),
        ("specialty_id", "Integer", "4 bytes, FK (specialties.id), NOT NULL", "Chuyên khoa công tác"),
        ("clinic_id", "Integer", "4 bytes, FK (clinics.id), NULLABLE", "Buồng khám phân công làm việc"),
        ("title", "String(50)", "Tối đa 50 ký tự, NULLABLE", "Học hàm, học vị: BS.CKI, ThS.BS, PGS.TS"),
        ("bio", "Text", "Tối đa 65,535 ký tự, NULLABLE", "Giới thiệu tóm tắt chuyên môn bác sĩ"),
        ("phone", "String(20)", "Tối đa 20 ký tự, NULLABLE", "Số điện thoại nội bộ phòng khám")
    ]
    tbl = doc.add_table(rows=len(doc_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(1.9)
    for idx, (c1, c2, c3, c4) in enumerate(doc_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp Doctor:")
    
    add_method_spec(
        doc,
        name="get_schedule",
        desc="Lấy danh sách ca trực và các lịch hẹn đã được đặt của bác sĩ trong một ngày cụ thể.",
        inputs="query_date: Date, kích thước 4 bytes (Ngày cần kiểm tra lịch).",
        output="schedule_list: List[Dictionary], chứa danh sách ca làm việc và các khoảng thời gian đã có hẹn.",
        flow="1. Xác định thứ trong tuần của query_date (day_of_week).\n2. Lọc danh sách Shift của bác sĩ có ngày trong tuần trùng khớp.\n3. Truy vấn các Appointment của bác sĩ trong ngày query_date có trạng thái khác CANCELLED.\n4. Trả về cấu trúc danh sách ca trực và lịch khám để hiển thị trên giao diện xếp lịch.",
        start_cond="query_date là ngày hợp lệ.",
        end_cond="Trả về toàn bộ dữ liệu lịch làm việc để hiển thị trên Calendar View."
    )

    add_method_spec(
        doc,
        name="is_available",
        desc="Kiểm tra bác sĩ có rảnh trong một khoảng thời gian cụ thể hay không.",
        inputs="start_time: DateTime, kích thước 8 bytes (Thời điểm bắt đầu ca khám dự kiến).\nend_time: DateTime, kích thước 8 bytes (Thời điểm kết thúc ca khám dự kiến).",
        output="available: Boolean, kích thước 1 byte (True nếu bác sĩ rảnh, False nếu đã bận lịch khác).",
        flow="1. Kiểm tra khoảng thời gian có nằm trong ca trực Shift của bác sĩ hay không.\n2. Kiểm tra xem có Appointment nào của bác sĩ đang trùng lặp thời gian hay không.\n3. Nếu không có lịch trùng, trả về True; ngược lại trả về False.",
        start_cond="start_time < end_time.",
        end_cond="Xác nhận tình trạng sẵn sàng tiếp nhận bệnh nhân của bác sĩ."
    )

    # -------------------------------------------------------------------------
    # 2.4 APPOINTMENT
    # -------------------------------------------------------------------------
    add_h2(doc, "2.4. Lớp Appointment (Lịch hẹn Khám bệnh)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Quản lý vòng đời lịch khám bệnh giữa bệnh nhân, bác sĩ và buồng khám, theo dõi trạng thái từ lúc đặt lịch đến khi hoàn tất khám.")

    add_h3(doc, "Các thuộc tính của lớp Appointment:")
    apt_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("id", "Integer", "4 bytes, PK, Auto Increment", "Định danh duy nhất lịch hẹn"),
        ("appointment_code", "String(30)", "Tối đa 30 ký tự, UNIQUE, NOT NULL, Index", "Mã lịch hẹn: LH-YYYYMMDD-XXXX"),
        ("patient_id", "Integer", "4 bytes, FK (patients.id), NOT NULL", "Bệnh nhân đặt lịch khám"),
        ("doctor_id", "Integer", "4 bytes, FK (doctors.id), NOT NULL", "Bác sĩ tiếp nhận ca khám"),
        ("clinic_id", "Integer", "4 bytes, FK (clinics.id), NULLABLE", "Buồng khám diễn ra ca khám"),
        ("appointment_date", "Date", "4 bytes, NOT NULL, Index", "Ngày hẹn khám bệnh"),
        ("start_time", "Time", "3 bytes, NOT NULL", "Giờ bắt đầu khám dự kiến"),
        ("end_time", "Time", "3 bytes, NOT NULL", "Giờ kết thúc khám dự kiến"),
        ("status", "String(20)", "Tối đa 20 ký tự, NOT NULL", "PENDING, CONFIRMED, CHECKED_IN, IN_PROGRESS, COMPLETED, CANCELLED"),
        ("reason", "String(255)", "Tối đa 255 ký tự, NULLABLE", "Lý do khám bệnh / Triệu chứng ban đầu"),
        ("notes", "Text", "Tối đa 65,535 ký tự, NULLABLE", "Ghi chú thêm của tiếp đón hoặc bác sĩ"),
        ("created_at", "DateTime", "8 bytes, DEFAULT utcnow, NOT NULL", "Thời điểm tạo lịch hẹn")
    ]
    tbl = doc.add_table(rows=len(apt_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(1.9)
    for idx, (c1, c2, c3, c4) in enumerate(apt_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp Appointment:")
    
    add_method_spec(
        doc,
        name="confirm_booking",
        desc="Lễ tân xác nhận lịch hẹn đã được chấp thuận sau khi kiểm tra tính hợp lệ.",
        inputs="Không có.",
        output="None, kích thước 0 byte.",
        flow="1. Kiểm tra trạng thái hiện tại: nếu self.status == 'PENDING'.\n2. Gán self.status = 'CONFIRMED'.",
        start_cond="Lịch hẹn ở trạng thái PENDING.",
        end_cond="Trạng thái chuyển thành CONFIRMED."
    )

    add_method_spec(
        doc,
        name="check_in",
        desc="Ghi nhận bệnh nhân đã có mặt tại phòng khám, chuyển vào danh sách chờ khám của bác sĩ.",
        inputs="Không có.",
        output="None, kích thước 0 byte.",
        flow="1. Kiểm tra trạng thái lịch hẹn phải là CONFIRMED.\n2. Gán self.status = 'CHECKED_IN'.",
        start_cond="Bệnh nhân đến quầy tiếp đón đúng ngày hẹn.",
        end_cond="Trạng thái chuyển thành CHECKED_IN, hiển thị trên hàng đợi của bác sĩ."
    )

    add_method_spec(
        doc,
        name="cancel_appointment",
        desc="Hủy lịch hẹn theo yêu cầu của bệnh nhân hoặc lý do đột xuất của bác sĩ.",
        inputs="cancel_reason: String, kích thước tối đa 255 ký tự (Lý do hủy lịch).",
        output="None, kích thước 0 byte.",
        flow="1. Kiểm tra trạng thái hiện tại: không cho phép hủy nếu ca khám đã IN_PROGRESS hoặc COMPLETED.\n2. Gán self.status = 'CANCELLED'.\n3. Cập nhật self.notes = f'{self.notes or \"\"} [Lý do hủy: {cancel_reason}]'.",
        start_cond="Lịch hẹn chưa bước vào ca khám bệnh.",
        end_cond="Trạng thái chuyển thành CANCELLED, giải phóng khung giờ trống cho bệnh nhân khác."
    )

    add_method_spec(
        doc,
        name="reschedule",
        desc="Đổi ngày hoặc giờ hẹn sang một khung giờ mới sau khi đã xác minh không bị xung đột lịch.",
        inputs="new_date: Date, kích thước 4 bytes (Ngày khám mới).\nnew_start: Time, kích thước 3 bytes (Giờ bắt đầu mới).\nnew_end: Time, kích thước 3 bytes (Giờ kết thúc mới).",
        output="success: Boolean, kích thước 1 byte (True nếu đổi lịch thành công).",
        flow="1. Cập nhật self.appointment_date = new_date.\n2. Cập nhật self.start_time = new_start.\n3. Cập nhật self.end_time = new_end.\n4. Đặt lại trạng thái self.status = 'CONFIRMED'.\n5. Trả về True.",
        start_cond="Khung giờ mới đã được kiểm tra qua ConflictChecker và không có xung đột.",
        end_cond="Thông tin ngày giờ lịch hẹn được cập nhật mới hoàn toàn."
    )

    # -------------------------------------------------------------------------
    # 2.5 CONFLICT CHECKER
    # -------------------------------------------------------------------------
    add_h2(doc, "2.5. Lớp ConflictChecker (Bộ Phát hiện Xung đột Lịch khám)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Tiện ích dịch vụ toán học kiểm tra sự giao thoa khoảng thời gian (Interval Intersection) trên hai chiều không gian: Lịch làm việc của Bác sĩ và Buồng khám của Phòng khám.")

    add_h3(doc, "Các thuộc tính của lớp ConflictChecker:")
    p_cc = doc.add_paragraph()
    p_cc.add_run("Lớp dịch vụ tiện ích độc lập (Stateless Service Class), không lưu trữ trạng thái trong CSDL.")

    add_h3(doc, "Các phương thức của lớp ConflictChecker:")
    add_method_spec(
        doc,
        name="check_appointment_conflict",
        desc="Kiểm tra yêu cầu đặt lịch mới hoặc đổi lịch có bị chồng lấn thời gian với bất kỳ lịch hẹn nào đang hoạt động hay không.",
        inputs="doctor_id: Integer, 4 bytes (Mã bác sĩ).\nclinic_id: Integer, 4 bytes, NULLABLE (Mã buồng khám).\napt_date: Date, 4 bytes (Ngày hẹn khám).\nstart_time: Time, 3 bytes (Giờ bắt đầu).\nend_time: Time, 3 bytes (Giờ kết thúc).\nexclude_id: Integer, 4 bytes, DEFAULT None (Mã lịch hẹn cần loại trừ khi cập nhật lịch cũ).",
        output="result: Tuple[Boolean, Optional[String]]\n- Phần tử 1: Boolean (True nếu KHÔNG xung đột, False nếu BỊ TRÙNG LỊCH).\n- Phần tử 2: String hoặc None (Thông báo chi tiết nguyên nhân xung đột).",
        flow="1. Ràng buộc thời gian: nếu start_time >= end_time, trả về (False, 'Giờ bắt đầu phải trước giờ kết thúc').\n2. Áp dụng công thức giao thoa: Overlap(A, B) <=> (Start_A < End_B) and (End_A > Start_B).\n3. Truy vấn bảng appointments tìm lịch của bác sĩ doctor_id tại ngày apt_date có trạng thái trong ['PENDING', 'CONFIRMED', 'CHECKED_IN', 'IN_PROGRESS']. Nếu tìm thấy và id != exclude_id, trả về (False, 'Bác sĩ đã có lịch hẹn khác trong khung giờ này').\n4. Nếu có clinic_id, kiểm tra buồng khám clinic_id tại ngày apt_date. Nếu có lịch hẹn khác đang sử dụng phòng, trả về (False, 'Phòng khám đã có ca khám khác trong khung giờ này').\n5. Nếu không phát hiện xung đột, trả về (True, None).",
        start_cond="Có đầy đủ thông tin bác sĩ, ngày và khung giờ cần kiểm tra.",
        end_cond="Trả về kết quả phân tích xung đột trong thời gian thực (< 5ms)."
    )

    add_method_spec(
        doc,
        name="check_shift_alignment",
        desc="Kiểm tra giờ khám bệnh yêu cầu có nằm trọn vẹn trong ca làm việc đăng ký của bác sĩ hay không.",
        inputs="doctor_id: Integer, 4 bytes.\napt_date: Date, 4 bytes.\nstart_time: Time, 3 bytes.\nend_time: Time, 3 bytes.",
        output="is_aligned: Boolean, kích thước 1 byte (True nếu nằm trong ca trực, False nếu ngoài giờ làm việc).",
        flow="1. Xác định thứ trong tuần của apt_date (day_of_week = apt_date.weekday()).\n2. Truy vấn bảng shifts tìm ca làm việc của doctor_id có day_of_week == day_of_week.\n3. Kiểm tra điều kiện: (shift.start_time <= start_time) and (shift.end_time >= end_time).\n4. Nếu thỏa mãn trả về True, ngược lại trả về False.",
        start_cond="Ngày khám và khung giờ hợp lệ.",
        end_cond="Ngăn chặn đặt lịch vào các khung giờ bác sĩ không có ca trực."
    )

    # -------------------------------------------------------------------------
    # 2.6 MEDICAL RECORD
    # -------------------------------------------------------------------------
    add_h2(doc, "2.6. Lớp MedicalRecord (Phiếu Khám bệnh Lâm sàng EMR)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Đại diện cho hồ sơ khám bệnh điện tử, lưu trữ thông tin đo sinh hiệu, triệu chứng, chẩn đoán ICD-10 của bác sĩ, đóng vai trò gốc liên kết tới các chỉ định cận lâm sàng, đơn thuốc và hóa đơn viện phí.")

    add_h3(doc, "Các thuộc tính của lớp MedicalRecord:")
    mr_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("id", "Integer", "4 bytes, PK, Auto Increment", "Định danh duy nhất phiếu khám"),
        ("record_code", "String(30)", "Tối đa 30 ký tự, UNIQUE, NOT NULL, Index", "Mã phiếu khám: KB-YYYYMMDD-XXXX"),
        ("patient_id", "Integer", "4 bytes, FK (patients.id), NOT NULL", "Bệnh nhân được khám"),
        ("doctor_id", "Integer", "4 bytes, FK (doctors.id), NOT NULL", "Bác sĩ thực hiện khám"),
        ("appointment_id", "Integer", "4 bytes, FK (appointments.id), NULLABLE", "Lịch hẹn liên quan"),
        ("exam_date", "DateTime", "8 bytes, DEFAULT utcnow, NOT NULL", "Thời điểm bắt đầu ca khám"),
        ("chief_complaint", "Text", "Tối đa 65,535 ký tự, NOT NULL", "Triệu chứng chính / Lý do đến khám"),
        ("blood_pressure", "String(20)", "Tối đa 20 ký tự, NULLABLE", "Huyết áp (e.g. 120/80 mmHg)"),
        ("heart_rate", "Integer", "4 bytes, NULLABLE", "Nhịp tim (lần/phút, bpm)"),
        ("temperature", "Float", "8 bytes, NULLABLE", "Thân nhiệt (°C)"),
        ("respiratory_rate", "Integer", "4 bytes, NULLABLE", "Nhịp thở (lần/phút)"),
        ("weight", "Float", "8 bytes, NULLABLE", "Cân nặng (kg)"),
        ("height", "Float", "8 bytes, NULLABLE", "Chiều cao (cm)"),
        ("bmi", "Float", "8 bytes, NULLABLE", "Chỉ số khối cơ thể (BMI)"),
        ("physical_exam", "Text", "Tối đa 65,535 ký tự, NULLABLE", "Khám lâm sàng các cơ quan"),
        ("diagnosis_icd10", "String(255)", "Tối đa 255 ký tự, NULLABLE", "Tên chẩn đoán bệnh chính"),
        ("icd10_code", "String(20)", "Tối đa 20 ký tự, NULLABLE", "Mã phân loại quốc tế ICD-10"),
        ("doctor_notes", "Text", "Tối đa 65,535 ký tự, NULLABLE", "Lời dặn dò chuyên môn"),
        ("status", "String(20)", "Tối đa 20 ký tự, DEFAULT 'IN_EXAM'", "Trạng thái: IN_EXAM, COMPLETED, CANCELLED"),
        ("created_at", "DateTime", "8 bytes, DEFAULT utcnow, NOT NULL", "Thời điểm lập phiếu khám")
    ]
    tbl = doc.add_table(rows=len(mr_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(1.9)
    for idx, (c1, c2, c3, c4) in enumerate(mr_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp MedicalRecord:")
    
    add_method_spec(
        doc,
        name="calculate_bmi",
        desc="Tự động tính toán chỉ số khối cơ thể (BMI) theo công thức chuẩn của Tổ chức Y tế Thế giới (WHO).",
        inputs="Không có.",
        output="bmi_value: Float, kích thước 8 bytes (Chỉ số BMI làm tròn 2 chữ số thập phân).",
        flow="1. Kiểm tra self.height và self.weight: nếu <= 0 hoặc None, trả về 0.0.\n2. Đổi chiều cao sang mét: h_m = self.height / 100.0.\n3. Áp dụng công thức: bmi = self.weight / (h_m * h_m).\n4. Làm tròn: self.bmi = round(bmi, 2).\n5. Trả về self.bmi.",
        start_cond="Thuộc tính weight và height đã được nhập liệu.",
        end_cond="Thuộc tính bmi của đối tượng được cập nhật giá trị chính xác."
    )

    add_method_spec(
        doc,
        name="add_service_order",
        desc="Bác sĩ chỉ định một dịch vụ cận lâm sàng (Xét nghiệm máu, chụp X-Quang, siêu âm).",
        inputs="service_code: String, tối đa 30 ký tự (Mã dịch vụ kỹ thuật, e.g. SA-BUNG-01).\nservice_name: String, tối đa 150 ký tự (Tên dịch vụ kỹ thuật).\nprice: Float, kích thước 8 bytes (Giá dịch vụ theo bảng giá niêm yết).",
        output="new_order: ServiceOrder, đối tượng chỉ định dịch vụ mới được khởi tạo.",
        flow="1. Khởi tạo đối tượng order = ServiceOrder(medical_record_id=self.id, service_code=service_code, service_name=service_name, price=price).\n2. Thêm vào danh sách self.service_orders.append(order).\n3. Trả về order.",
        start_cond="Phiếu khám đang ở trạng thái IN_EXAM.",
        end_cond="Bản ghi chỉ định dịch vụ được gắn kết với phiếu khám bệnh."
    )

    add_method_spec(
        doc,
        name="create_prescription",
        desc="Khởi tạo đơn thuốc cho ca khám bệnh hiện tại.",
        inputs="advice: String, tối đa 65,535 ký tự (Lời dặn dò sử dụng thuốc của bác sĩ).",
        output="prescription: Prescription, đối tượng đơn thuốc liên kết 1-1 với phiếu khám.",
        flow="1. Sinh mã đơn thuốc duy nhất: code = f'DT-{datetime.utcnow().strftime(\"%Y%m%d\")}-{uuid.uuid4().hex[:4].upper()}'.\n2. Khởi tạo đối tượng Prescription(prescription_code=code, medical_record_id=self.id, doctor_id=self.doctor_id, patient_id=self.patient_id, diagnosis=self.diagnosis_icd10, advice=advice).\n3. Gán self.prescription = rx.\n4. Trả về rx.",
        start_cond="Phiếu khám đã có chẩn đoán bệnh diagnosis_icd10.",
        end_cond="Đơn thuốc được tạo và liên kết trực tiếp với phiếu khám."
    )

    add_method_spec(
        doc,
        name="complete_examination",
        desc="Bác sĩ hoàn thành ca khám lâm sàng, khóa bệnh án và chuyển hồ sơ sang bộ phận thu ngân kế toán để lập hóa đơn viện phí.",
        inputs="Không có.",
        output="None, kích thước 0 byte.",
        flow="1. Kiểm tra chẩn đoán: bắt buộc self.diagnosis_icd10 không được để trống.\n2. Gán trạng thái self.status = 'COMPLETED'.\n3. Nếu có liên kết với Appointment, cập nhật self.appointment.status = 'COMPLETED'.",
        start_cond="Bác sĩ đã hoàn tất chẩn đoán lâm sàng và ký đơn thuốc/chỉ định.",
        end_cond="Trạng thái phiếu khám chuyển thành COMPLETED, sẵn sàng cho thanh toán viện phí."
    )

    # -------------------------------------------------------------------------
    # 2.7 MEDICINE & PRESCRIPTION
    # -------------------------------------------------------------------------
    add_h2(doc, "2.7. Lớp Medicine (Danh mục Thuốc & Quản lý Kho Dược)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Quản lý kho dược phẩm phòng khám, theo dõi số lượng tồn kho, hoạt chất, đơn giá và cảnh báo cạn kho.")

    add_h3(doc, "Các thuộc tính của lớp Medicine:")
    med_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("id", "Integer", "4 bytes, PK, Auto Increment", "Định danh duy nhất của thuốc"),
        ("code", "String(30)", "Tối đa 30 ký tự, UNIQUE, NOT NULL, Index", "Mã thuốc: MED-PARA-500"),
        ("name", "String(150)", "Tối đa 150 ký tự, NOT NULL, Index", "Tên thương mại của thuốc"),
        ("active_ingredient", "String(150)", "Tối đa 150 ký tự, NOT NULL", "Tên hoạt chất chính"),
        ("dosage_form", "String(50)", "Tối đa 50 ký tự, NOT NULL", "Dạng bào chế: Viên nén, Siro, Viên nang"),
        ("unit", "String(30)", "Tối đa 30 ký tự, NOT NULL", "Đơn vị tính: Viên, Vỉ, Hộp, Chai"),
        ("unit_price", "Float", "8 bytes, DEFAULT 0.0, NOT NULL", "Đơn giá bán lẻ (VNĐ)"),
        ("stock_quantity", "Integer", "4 bytes, DEFAULT 0, NOT NULL", "Số lượng còn lại trong kho"),
        ("usage_instructions", "Text", "Tối đa 65,535 ký tự, NULLABLE", "Hướng dẫn sử dụng mặc định"),
        ("is_active", "Boolean", "1 byte, DEFAULT True, NOT NULL", "Trạng thái lưu hành thuốc")
    ]
    tbl = doc.add_table(rows=len(med_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(1.9)
    for idx, (c1, c2, c3, c4) in enumerate(med_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp Medicine:")
    
    add_method_spec(
        doc,
        name="deduct_stock",
        desc="Trừ số lượng tồn kho của thuốc khi xuất kho cấp phát thuốc theo đơn.",
        inputs="qty: Integer, kích thước 4 bytes (Số lượng thuốc cần xuất kho).",
        output="success: Boolean, kích thước 1 byte (True nếu xuất kho thành công, False nếu tồn kho không đủ).",
        flow="1. Kiểm tra tồn kho: nếu self.stock_quantity < qty, trả về False.\n2. Thực hiện trừ kho: self.stock_quantity -= qty.\n3. Trả về True.",
        start_cond="qty > 0.",
        end_cond="Tồn kho thuốc được giảm tương ứng với số lượng xuất."
    )

    add_method_spec(
        doc,
        name="is_in_stock",
        desc="Kiểm tra kho dược có đủ số lượng đáp ứng đơn thuốc yêu cầu hay không.",
        inputs="required_qty: Integer, kích thước 4 bytes (Số lượng cần kiểm tra).",
        output="in_stock: Boolean, kích thước 1 byte (True nếu stock_quantity >= required_qty).",
        flow="1. Trả về self.stock_quantity >= required_qty and self.is_active.",
        start_cond="Có số lượng yêu cầu cần đối soát.",
        end_cond="Ngăn chặn việc bác sĩ kê đơn thuốc đã cạn kiệt trong kho."
    )

    # -------------------------------------------------------------------------
    # 2.8 INVOICE
    # -------------------------------------------------------------------------
    add_h2(doc, "2.8. Lớp Invoice (Hóa đơn Viện phí & Thanh toán VietQR)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Tổng hợp các khoản chi phí của ca khám (công khám bệnh, phí xét nghiệm cận lâm sàng, tiền thuốc), tự động tính toán khấu trừ BHYT và sinh mã thanh toán VietQR động.")

    add_h3(doc, "Các thuộc tính của lớp Invoice:")
    inv_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("id", "Integer", "4 bytes, PK, Auto Increment", "Định danh duy nhất của hóa đơn"),
        ("invoice_code", "String(30)", "Tối đa 30 ký tự, UNIQUE, NOT NULL, Index", "Mã hóa đơn viện phí: HD-YYYYMMDD-XXXX"),
        ("medical_record_id", "Integer", "4 bytes, FK (medical_records.id), NULLABLE", "Ca khám bệnh phát sinh chi phí"),
        ("patient_id", "Integer", "4 bytes, FK (patients.id), NOT NULL", "Bệnh nhân chi trả"),
        ("consultation_fee", "Float", "8 bytes, DEFAULT 150000.0, NOT NULL", "Tiền công khám bác sĩ (VNĐ)"),
        ("service_fee", "Float", "8 bytes, DEFAULT 0.0, NOT NULL", "Tiền các dịch vụ cận lâm sàng (VNĐ)"),
        ("medicine_fee", "Float", "8 bytes, DEFAULT 0.0, NOT NULL", "Tổng tiền thuốc kê đơn (VNĐ)"),
        ("total_amount", "Float", "8 bytes, DEFAULT 0.0, NOT NULL", "Tổng chi phí trước giảm trừ (VNĐ)"),
        ("insurance_discount", "Float", "8 bytes, DEFAULT 0.0, NOT NULL", "Số tiền BHYT chi trả giảm trừ (VNĐ)"),
        ("patient_pay_amount", "Float", "8 bytes, DEFAULT 0.0, NOT NULL", "Số tiền thực tế người bệnh phải nộp (VNĐ)"),
        ("payment_status", "String(20)", "Tối đa 20 ký tự, DEFAULT 'PENDING'", "Trạng thái: PENDING, PAID, CANCELLED"),
        ("payment_method", "String(20)", "Tối đa 20 ký tự, NULLABLE", "Hình thức: CASH, BANK_TRANSFER, INSURANCE"),
        ("transaction_code", "String(50)", "Tối đa 50 ký tự, NULLABLE", "Mã giao dịch chuyển khoản / VietQR ref"),
        ("cashier_id", "Integer", "4 bytes, FK (users.id), NULLABLE", "Nhân viên thu ngân tiếp nhận thanh toán"),
        ("paid_at", "DateTime", "8 bytes, NULLABLE", "Thời điểm thanh toán thành công"),
        ("created_at", "DateTime", "8 bytes, DEFAULT utcnow, NOT NULL", "Thời điểm lập hóa đơn")
    ]
    tbl = doc.add_table(rows=len(inv_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(1.9)
    for idx, (c1, c2, c3, c4) in enumerate(inv_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp Invoice:")
    
    add_method_spec(
        doc,
        name="calculate_total_bill",
        desc="Tổng hợp các khoản viện phí và tính toán số tiền đồng chi trả sau khi áp dụng mức hưởng BHYT.",
        inputs="bhyt_rate: Float, kích thước 8 bytes (Tỷ lệ chi trả của BHYT từ 0.0 đến 1.0, ví dụ 0.8 cho 80% đúng tuyến hoặc 1.0 cho 100%).",
        output="None, kích thước 0 byte.",
        flow="1. Tính tổng chi phí trước giảm trừ: self.total_amount = self.consultation_fee + self.service_fee + self.medicine_fee.\n2. Tính mức BHYT giảm trừ: self.insurance_discount = (self.consultation_fee + self.service_fee) * bhyt_rate.\n3. Tính số tiền người bệnh thực trả: self.patient_pay_amount = max(0.0, self.total_amount - self.insurance_discount).",
        start_cond="Các khoản chi phí công khám, cận lâm sàng, thuốc đã được tổng hợp từ ca khám.",
        end_cond="Thuộc tính total_amount, insurance_discount và patient_pay_amount được cập nhật chính xác."
    )

    add_method_spec(
        doc,
        name="generate_vietqr_payload",
        desc="Tạo dữ liệu chuẩn để sinh mã QR thanh toán nhanh Napas247 theo định dạng chuẩn VietQR động.",
        inputs="Không có.",
        output="payload: Dictionary, chứa thông tin ngân hàng thụ hưởng, số tài khoản, số tiền phải trả và nội dung chuyển khoản.",
        flow="1. Xác định số tiền chuyển khoản: amount = int(self.patient_pay_amount).\n2. Tạo nội dung chuyển khoản chuẩn hóa: memo = f'CLINIC {self.invoice_code}'.\n3. Tạo URL mã QR thanh toán VietQR động: qr_url = f'https://img.vietqr.io/image/MB-0348736868-compact2.png?amount={amount}&addInfo={memo}&accountName=PHONG%20KHAM%20DA%20KHOA'.\n4. Trả về Dictionary chứa thông tin thanh toán.",
        start_cond="self.patient_pay_amount > 0 và hóa đơn đang ở trạng thái PENDING.",
        end_cond="Trả về dữ liệu để hiển thị mã QR trên màn hình thu ngân và in phiếu thu cho bệnh nhân."
    )

    add_method_spec(
        doc,
        name="process_payment",
        desc="Ghi nhận giao dịch thanh toán thành công từ thu ngân, cập nhật thời gian và khóa hóa đơn.",
        inputs="method: String, tối đa 20 ký tự (CASH, BANK_TRANSFER, INSURANCE).\ntx_code: String, tối đa 50 ký tự (Mã giao dịch ngân hàng hoặc số biên lai).\ncashier_user_id: Integer, kích thước 4 bytes (ID nhân viên thu ngân tiếp nhận tiền).",
        output="success: Boolean, kích thước 1 byte (True nếu ghi nhận thanh toán thành công).",
        flow="1. Kiểm tra trạng thái: nếu self.payment_status == 'PAID', trả về False (Tránh thanh toán trùng).\n2. Cập nhật self.payment_status = 'PAID'.\n3. Cập nhật self.payment_method = method.\n4. Cập nhật self.transaction_code = tx_code.\n5. Cập nhật self.cashier_id = cashier_user_id.\n6. Ghi nhận thời gian self.paid_at = datetime.utcnow().\n7. Trả về True.",
        start_cond="Hóa đơn đang ở trạng thái PENDING.",
        end_cond="Trạng thái chuyển thành PAID, giao dịch được ghi nhận vĩnh viễn vào hệ thống kế toán."
    )

    # -------------------------------------------------------------------------
    # 2.9 PII ANONYMIZER & AI ENGINE
    # -------------------------------------------------------------------------
    add_h2(doc, "2.9. Lớp PIIAnonymizer (Module Khử Định danh Dữ liệu Y tế)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Chịu trách nhiệm bảo vệ quyền riêng tư dữ liệu bệnh nhân (theo Nghị định 13/2023/NĐ-CP), tự động phát hiện và che giấu toàn bộ các thông tin định danh cá nhân nhạy cảm trước khi gửi prompt tới AI Engine.")

    add_h3(doc, "Các thuộc tính của lớp PIIAnonymizer:")
    p_pii = doc.add_paragraph()
    p_pii.add_run("Lớp tiện ích xử lý chuỗi và biểu thức chính quy (Regex Utility Class), cấu hình tập mẫu nhận diện PII (CCCD, SĐT, BHYT).")

    add_h3(doc, "Các phương thức của lớp PIIAnonymizer:")
    add_method_spec(
        doc,
        name="anonymize",
        desc="Thay thế các thông tin nhạy cảm (Số CCCD, Số điện thoại, Mã BHYT, Họ tên bệnh nhân) bằng các thẻ ẩn danh [REDACTED_...].",
        inputs="text: String, chuỗi văn bản hồ sơ bệnh án ban đầu cần khử định danh.",
        output="result: Tuple[String, Dictionary]\n- Phần tử 1: String, chuỗi văn bản đã được khử định danh 100% PII.\n- Phần tử 2: Dictionary, bảng ánh xạ lưu trữ tạm thời trong bộ nhớ đệm để khôi phục khi cần (mask_mapping).",
        flow="1. Áp dụng Regex phát hiện số CCCD (12 chữ số) -> Thay bằng [CCCD_REDACTED].\n2. Áp dụng Regex phát hiện số điện thoại Việt Nam (10 chữ số) -> Thay bằng [PHONE_REDACTED].\n3. Áp dụng Regex phát hiện mã thẻ BHYT (15 ký tự chữ và số) -> Thay bằng [BHYT_REDACTED].\n4. Lưu lại bảng ánh xạ giá trị thực và thẻ thay thế.\n5. Trả về chuỗi văn bản an toàn và bảng ánh xạ.",
        start_cond="Chuỗi văn bản đầu vào không rỗng.",
        end_cond="Dữ liệu được bảo vệ an toàn, không còn thông tin PII nhạy cảm trước khi rời khỏi máy chủ nội bộ."
    )

    add_method_spec(
        doc,
        name="deanonymize",
        desc="Khôi phục lại các thông tin ban đầu từ thẻ ẩn danh sau khi AI sinh phản hồi (nếu cần hiển thị trên giao diện bác sĩ).",
        inputs="text: String, chuỗi văn bản do AI trả về có chứa các thẻ ẩn danh.\nmask_mapping: Dictionary, bảng ánh xạ đã lưu ở bước khử định danh.",
        output="restored_text: String, chuỗi văn bản đã được khôi phục đầy đủ.",
        flow="1. Duyệt qua từng cặp (token, original_value) trong mask_mapping.\n2. Thay thế token trong văn bản bằng original_value.\n3. Trả về chuỗi văn bản hoàn chỉnh.",
        start_cond="Có chuỗi văn bản và bảng ánh xạ tương ứng.",
        end_cond="Khôi phục đúng ngữ cảnh hiển thị cho nhân viên y tế nội bộ."
    )

    # -------------------------------------------------------------------------
    # 2.10 ADMIN AI SERVICE
    # -------------------------------------------------------------------------
    add_h2(doc, "2.10. Lớp AdminAIService (Trợ lý AI Hành chính 3 Lớp)")
    p = doc.add_paragraph()
    p.add_run("Mô tả: ").font.bold = True
    p.add_run("Cung cấp 3 tính năng AI hành chính cốt lõi phục vụ bác sĩ, bệnh nhân và phòng khám, tích hợp đa nhà cung cấp (Google Gemini Live API và Offline Mock Fallback Engine).")

    add_h3(doc, "Các thuộc tính của lớp AdminAIService:")
    ai_attrs = [
        ("Tên thuộc tính", "Kiểu dữ liệu", "Kích thước / Ràng buộc", "Mô tả ý nghĩa"),
        ("anonymizer", "PIIAnonymizer", "Đối tượng tham chiếu", "Module thực hiện khử định danh PII"),
        ("guardrails", "AdminAIGuardrails", "Đối tượng tham chiếu", "Bộ lọc đạo đức y tế và bảo vệ prompt"),
        ("provider", "AIProvider", "Đối tượng tham chiếu", "Nhà cung cấp LLM (Gemini hoặc Offline Mock)")
    ]
    tbl = doc.add_table(rows=len(ai_attrs), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.3)
    tbl.columns[2].width = Inches(1.8)
    tbl.columns[3].width = Inches(2.1)
    for idx, (c1, c2, c3, c4) in enumerate(ai_attrs):
        row = tbl.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_h3(doc, "Các phương thức của lớp AdminAIService:")
    
    add_method_spec(
        doc,
        name="generate_pre_visit_summary",
        desc="Tóm tắt tiền sử bệnh án, các lần khám trước và làm nổi bật các cảnh báo dị ứng thuốc nguy hiểm để bác sĩ nắm bắt nhanh trong 10 giây trước khi vào ca khám.",
        inputs="patient_id: Integer, kích thước 4 bytes (Mã định danh bệnh nhân).",
        output="summary_result: Dictionary, chứa các mục: tiền sử mạn tính, tóm tắt các lần khám gần nhất, cảnh báo dị ứng thuốc nổi bật và tuyên bố miễn trừ y tế.",
        flow="1. Truy vấn thông tin bệnh nhân và tối đa 5 phiếu khám gần nhất của bệnh nhân từ CSDL.\n2. Trích xuất các trường: chẩn đoán cũ, thuốc đã dùng, ghi chú dị ứng.\n3. Khử định danh PII toàn bộ hồ sơ qua self.anonymizer.anonymize().\n4. Đóng gói prompt tóm tắt gửi đến self.provider.generate().\n5. Định dạng kết quả qua self.guardrails.format_output().\n6. Ghi nhật ký vào bảng ai_invocation_logs.\n7. Trả về kết quả Dictionary.",
        start_cond="Bác sĩ mở giao diện phòng khám lâm sàng đối với một bệnh nhân cụ thể.",
        end_cond="Trả về thẻ tóm tắt AI trên màn hình EMR trước khi tiến hành thăm khám."
    )

    add_method_spec(
        doc,
        name="answer_clinic_faq",
        desc="Trả lời tự động các câu hỏi của người bệnh về quy trình khám bệnh, thủ tục BHYT, giờ làm việc và bảng giá; từ chối mọi yêu cầu chẩn đoán bệnh tật.",
        inputs="question: String, nội dung câu hỏi của người bệnh bằng ngôn ngữ tự nhiên.",
        output="faq_response: Dictionary, chứa câu trả lời bằng tiếng Việt thân thiện, danh mục nguồn tham khảo và nhãn miễn trừ y tế.",
        flow="1. Kiểm tra an toàn qua self.guardrails.is_medical_diagnosis_query(question).\n2. Nếu phát hiện hỏi về chẩn đoán hoặc kê đơn, lập tức trả về lời từ chối lịch sự: 'Trợ lý AI chỉ hỗ trợ giải đáp quy trình hành chính phòng khám, không có chức năng chẩn đoán bệnh. Vui lòng đặt lịch khám để được bác sĩ tư vấn trực tiếp.'\n3. Nếu là câu hỏi hành chính, tra cứu tri thức trong RAG Knowledge Base của phòng khám (giờ mở cửa, bảng giá, thủ tục thẻ BHYT).\n4. Tạo prompt ghép dữ liệu ngữ cảnh gửi đến mô hình LLM.\n5. Đính kèm Medical Disclaimer và trả về cho người dùng.",
        start_cond="Người dùng gửi câu hỏi từ cửa sổ Chatbot.",
        end_cond="Trả về câu trả lời chính xác, an toàn và đúng quy chuẩn phòng khám."
    )

    add_method_spec(
        doc,
        name="generate_discharge_instructions",
        desc="Tự động sinh nội dung dặn dò sinh hoạt sau khám, phân chia lịch uống thuốc (Sáng - Trưa - Chiều - Tối), chế độ ăn uống kiêng cữ, dấu hiệu cấp cứu cần tái khám ngay.",
        inputs="medical_record_id: Integer, kích thước 4 bytes (Mã phiếu khám bệnh đã hoàn thành).",
        output="instructions_result: Dictionary, gồm bảng chia lịch uống thuốc, chế độ dinh dưỡng, dấu hiệu cảnh báo đỏ và lịch hẹn tái khám đề xuất.",
        flow="1. Truy vấn phiếu khám medical_record_id kèm đơn thuốc prescription liên quan.\n2. Trích xuất chẩn đoán và danh sách các thuốc kèm liều lượng, hướng dẫn uống.\n3. Tạo cấu trúc prompt yêu cầu AI định dạng bảng lịch uống thuốc trực quan, dễ hiểu cho người cao tuổi.\n4. Gọi AI sinh nội dung và gắn cảnh báo y tế.\n5. Chuyển cho bác sĩ xem xét và phê duyệt trên giao diện trước khi in gửi bệnh nhân.",
        start_cond="Phiếu khám đã hoàn thành và có đơn thuốc hợp lệ.",
        end_cond="Sinh tờ hướng dẫn dặn dò xuất viện hoàn chỉnh để in kèm đơn thuốc."
    )

    # =========================================================================
    # MỤC 3: BẢNG TỔNG HỢP & MA TRẬN ÁNH XẠ ĐÁP ỨNG QUY TRÌNH HƯỚNG ĐỐI TƯỢNG
    # =========================================================================
    add_h1(doc, "3. BẢNG TỔNG HỢP & MA TRẬN ÁNH XẠ THIẾT KẾ HƯỚNG ĐỐI TƯỢNG")
    p = doc.add_paragraph()
    p.add_run("Bảng ma trận phân tầng trách nhiệm và mối liên kết giữa các lớp nghiệp vụ trong hệ thống:")
    
    matrix_data = [
        ("Tên Lớp", "Trách nhiệm chính trong hệ thống", "Phân tầng kiến trúc", "Mối quan hệ chính"),
        ("User", "Xác thực JWT, phân quyền 4 vai trò RBAC, băm mật khẩu bcrypt", "Persistence & Security", "User 1-1 Doctor"),
        ("Patient", "Quản lý hồ sơ nhân khẩu, tiền sử bệnh, dị ứng thuốc", "Domain Entity", "Patient 1-N Appointment, Patient 1-N MedicalRecord"),
        ("Doctor", "Quản lý bác sĩ chuyên khoa, buồng khám và lịch trực", "Domain Entity", "Doctor 1-N Shift, Doctor 1-N Appointment"),
        ("Appointment", "Quản lý lịch hẹn khám, kiểm soát trạng thái đặt lịch", "Domain Entity", "Appointment 1-1 MedicalRecord"),
        ("ConflictChecker", "Thuật toán kiểm tra trùng lịch khám theo Interval Overlap", "Core Service", "Dependency tới Appointment & Shift"),
        ("MedicalRecord", "Bệnh án điện tử EMR, sinh hiệu, chẩn đoán ICD-10", "Domain Entity", "MedicalRecord 1-N ServiceOrder, 1-1 Prescription"),
        ("Prescription", "Đơn thuốc điều trị, lời dặn dùng thuốc của bác sĩ", "Domain Entity", "Prescription 1-N PrescriptionItem"),
        ("Medicine", "Quản lý kho dược phẩm, theo dõi tồn kho và xuất nhập thuốc", "Domain Entity", "Medicine 1-N PrescriptionItem"),
        ("Invoice", "Viện phí, khấu trừ BHYT, sinh mã thanh toán VietQR", "Financial Entity", "Invoice 1-1 MedicalRecord"),
        ("PIIAnonymizer", "Khử định danh PII (CCCD, SĐT, BHYT) theo Nghị định 13", "AI Security Utility", "Dependency tới AdminAIService"),
        ("AdminAIGuardrails", "Kiểm soát đạo đức y tế, chặn câu hỏi tự chẩn đoán", "AI Safety Guardrail", "Dependency tới AdminAIService"),
        ("AdminAIService", "Cung cấp 3 tính năng AI tóm tắt, hỏi đáp và dặn dò sau khám", "AI Application Service", "Tích hợp Gemini API & Mock Fallback Engine")
    ]
    tbl_m = doc.add_table(rows=len(matrix_data), cols=4)
    tbl_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_m.columns[0].width = Inches(1.3)
    tbl_m.columns[1].width = Inches(2.2)
    tbl_m.columns[2].width = Inches(1.5)
    tbl_m.columns[3].width = Inches(1.5)
    for idx, (c1, c2, c3, c4) in enumerate(matrix_data):
        row = tbl_m.rows[idx]
        row.cells[0].text, row.cells[1].text, row.cells[2].text, row.cells[3].text = c1, c2, c3, c4
        if idx == 0:
            format_row(row, "0D9488", RGBColor(255, 255, 255), is_header=True)
        else:
            format_row(row, "FFFFFF" if idx % 2 == 1 else "F8FAFC", RGBColor(30, 41, 59))

    # =========================================================================
    # KÝ TÊN VÀ XÁC NHẬN CỦA NHÓM 07
    # =========================================================================
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.columns[0].width = Inches(3.2)
    sign_table.columns[1].width = Inches(3.3)
    
    c_left = sign_table.cell(0, 0)
    c_right = sign_table.cell(0, 1)
    
    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.add_run("THÀNH VIÊN NHÓM 07\n(Ký và ghi rõ họ tên)\n\n\n\n").font.bold = True
    p_l.add_run("Trần Đặng Công Tâm\n(Fullstack Dev & QA Lead)").font.italic = True
    
    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.add_run("TRƯỞNG NHÓM 07\n(Ký và ghi rõ họ tên)\n\n\n\n").font.bold = True
    p_r.add_run("Đinh Gia Bảo\n(Architecture & AI Lead)").font.italic = True

    # Lưu tài liệu vào các thư mục docs/ và docs/Reports/
    out_paths = [
        r"d:\ICTU\Nam 3\ICTU_2026-2027\Ứng dụng trí tuệ nhân tạo - Project\He_thong_quan_ly_phong_kham\docs\Tai_Lieu_Thiet_Ke_Huong_Doi_Tuong_Nhom_07.docx",
        r"d:\ICTU\Nam 3\ICTU_2026-2027\Ứng dụng trí tuệ nhân tạo - Project\He_thong_quan_ly_phong_kham\docs\Reports\Tai_Lieu_Thiet_Ke_Huong_Doi_Tuong_Nhom_07.docx",
        r"d:\ICTU\Nam 3\ICTU_2026-2027\Ứng dụng trí tuệ nhân tạo - Project\He_thong_quan_ly_phong_kham\Tai_Lieu_Thiet_Ke_Huong_Doi_Tuong_Nhom_07.docx"
    ]
    
    for p in out_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        doc.save(p)
    print("Successfully generated all docx files!")

if __name__ == "__main__":
    build_docx()
