# -*- coding: utf-8 -*-
"""
Script sinh file 07_GenAI_SoftwareDevelopment_user-guide.docx
Sổ tay Hướng dẫn sử dụng hệ thống CMS-AI cho 4 vai trò (Nhóm 03).
Bổ sung:
- Sơ đồ tương tác luồng giao diện 4 vai trò (diagram_ui_flows_role.png)
- Hướng dẫn thao tác chi tiết từng trường nhập liệu cho Lễ tân, Bác sĩ, Kế toán, Admin
- Bảng tiêu chí chấm điểm và phiếu đánh giá kịch bản Demo 8 bước cho Hội đồng nghiệm thu
- Sổ tay chẩn đoán và khắc phục sự cố kỹ thuật thường gặp (Troubleshooting Guide)
"""

import os
import sys
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_style_helpers import (
    create_base_document,
    add_header_block,
    add_team_meta,
    add_h1,
    add_h2,
    add_h3,
    add_p,
    add_diagram,
    add_code_block,
    format_row,
    set_table_borders,
    save_document,
    HEX_HEADER_BG,
    HEX_ROW_ALT
)

def build_user_guide():
    doc = create_base_document()

    add_header_block(
        doc,
        "HƯỚNG DẪN SỬ DỤNG HỆ THỐNG (USER GUIDE)",
        "HỆ THỐNG QUẢN LÝ PHÒNG KHÁM CÓ TÍCH HỢP AI (CMS-AI)"
    )

    add_team_meta(doc)

    # 1. GIỚI THIỆU ỨNG DỤNG
    add_h1(doc, "1. GIỚI THIỆU ỨNG DỤNG")
    add_p(
        doc,
        "Hệ thống quản lý phòng khám đa khoa thông minh CMS-AI là giải pháp phần mềm chuyên biệt phục vụ các cơ sở y tế tư nhân quy mô vừa và nhỏ. "
        "Ứng dụng số hóa toàn diện quy trình tiếp đón người bệnh, điều phối lịch hẹn, quản lý hồ sơ bệnh án điện tử EMR, "
        "kê đơn thuốc điện tử và thanh toán viện phí không dùng tiền mặt qua mã VietQR chuẩn quốc gia. "
        "Điểm đột phá của hệ thống nằm ở động cơ Trợ lý AI Hành chính 3 lớp (AI Assistant): giúp bác sĩ tóm tắt hồ sơ cũ trong 15 giây, "
        "hỗ trợ người bệnh hỏi đáp quy trình khám tự động (RAG) và tự động sinh bản dặn dò sau khám theo mẫu y tế chuẩn mực."
    )

    # 2. CẤU HÌNH PHẦN CỨNG - PHẦN MỀM
    add_h1(doc, "2. CẤU HÌNH PHẦN CỨNG - PHẦN MỀM")

    add_h2(doc, "2.1. Cấu hình phần cứng")
    add_p(
        doc,
        "Máy chủ (Server): Bộ vi xử lý Intel Core i5/i7 thế hệ 11 trở lên hoặc AMD Ryzen 5/7; Bộ nhớ RAM tối thiểu 8 GB (khuyến nghị 16 GB); "
        "Ổ cứng thể rắn SSD dung lượng trống tối thiểu 20 GB; Kết nối mạng LAN hoặc Internet tốc độ cao.\n"
        "Máy trạm nhân viên (Client): Máy tính để bàn, máy tính xách tay hoặc máy tính bảng kết nối cùng mạng nội bộ; "
        "Màn hình độ phân giải tối thiểu 1366x768 pixels."
    )

    add_h2(doc, "2.2. Cấu hình phần mềm")
    add_p(
        doc,
        "Phía máy chủ: Hệ điều hành Windows 10/11 Pro hoặc Ubuntu Linux 22.04 LTS; "
        "Nền tảng ảo hóa Docker Desktop 4.x và Docker Compose; Trình quản lý gói Python 3.11/3.12.\n"
        "Phía máy trạm: Trình duyệt web hiện đại phiên bản mới nhất như Google Chrome, Microsoft Edge hoặc Mozilla Firefox."
    )

    add_h2(doc, "2.3. Hướng dẫn khởi động nhanh hệ thống")
    add_p(
        doc,
        "Người quản trị có thể khởi động toàn bộ hệ thống bằng một trong hai phương thức tiện lợi sau:"
    )

    startup_guide_ascii = """+----------------------------------------------------------------------------------------------------+
|                                 HƯỚNG DẪN KHỞI CHẠY HỆ THỐNG PHÒNG KHÁM                            |
+----------------------------------------------------------------------------------------------------+
  Cách 1: Khởi chạy 1-Click trên Windows (Khuyến nghị cho phòng khám)
  - Bước 1: Mở thư mục chứa mã nguồn dự án.
  - Bước 2: Nhấp đúp chuột vào tệp tin 'run_app.bat' hoặc 'start_demo.bat'.
  - Bước 3: Hệ thống tự động kích hoạt CSDL MySQL/PostgreSQL, Backend FastAPI và Frontend React.
  - Bước 4: Trình duyệt web tự động mở địa chỉ giao diện: http://localhost:5173

  Cách 2: Khởi chạy qua Docker Compose (Môi trường máy chủ chuyên nghiệp)
  - Lệnh khởi động: docker-compose up -d --build
  - Địa chỉ Frontend: http://localhost:3000 (hoặc http://localhost:5173)
  - Địa chỉ Backend Swagger API: http://localhost:8000/docs
  - Địa chỉ Quản trị CSDL: http://localhost:8080
+----------------------------------------------------------------------------------------------------+"""
    add_code_block(doc, startup_guide_ascii)

    add_h2(doc, "2.4. Danh mục tài khoản thử nghiệm theo vai trò")
    tbl_acc = doc.add_table(rows=5, cols=4)
    tbl_acc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_acc.autofit = False
    col_w_acc = [Inches(1.8), Inches(1.5), Inches(1.4), Inches(1.8)]
    for row in tbl_acc.rows:
        for idx, w in enumerate(col_w_acc):
            row.cells[idx].width = w

    headers_acc = ["Vai trò người dùng", "Tên đăng nhập", "Mật khẩu mặc định", "Phạm vi quyền hạn"]
    for i, h in enumerate(headers_acc):
        tbl_acc.cell(0, i).text = h
    format_row(tbl_acc.rows[0], HEX_HEADER_BG, RGBColor(255, 255, 255), bold=True, is_header=True)

    acc_data = [
        ("Quản trị viên (Admin)", "admin", "Admin@123", "Toàn quyền quản trị hệ thống, ca trực, kho dược"),
        ("Lễ tân (Receptionist)", "receptionist", "Recep@123", "Tiếp đón người bệnh, đặt lịch hẹn, điều phối hàng đợi"),
        ("Bác sĩ (Doctor)", "doctor", "Doctor@123", "Bàn khám EMR, xem tóm tắt AI, chẩn đoán ICD-10, kê đơn"),
        ("Kế toán (Accountant)", "accountant", "Acc@123", "Quản lý hóa đơn viện phí, BHYT, thanh toán VietQR")
    ]
    for r_idx, (a_role, a_user, a_pass, a_perm) in enumerate(acc_data):
        row = tbl_acc.rows[r_idx + 1]
        row.cells[0].text = a_role
        row.cells[1].text = a_user
        row.cells[2].text = a_pass
        row.cells[3].text = a_perm
        bg = HEX_ROW_ALT if (r_idx % 2 == 1) else "FFFFFF"
        format_row(row, bg, RGBColor(15, 23, 42), bold=False, is_header=False)
    set_table_borders(tbl_acc)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. CÁC CHỨC NĂNG CHÍNH THEO 4 VAI TRÒ
    add_h1(doc, "3. CÁC CHỨC NĂNG CHÍNH VÀ LUỒNG THAO TÁC GIAO DIỆN")
    add_p(
        doc,
        "Hệ thống CMS-AI thiết kế luồng màn hình trực quan, công thái học, giúp từng bộ phận nhân sự "
        "thực hiện công việc chuyên trách một cách nhanh chóng và chính xác:"
    )

    # Chèn sơ đồ màn hình 4 vai trò
    flow_role_img = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "generated_diagrams", "diagram_ui_flows_role.png"))
    add_diagram(doc, flow_role_img, "Hình 3.1: Sơ đồ luồng giao diện người dùng theo 4 vai trò chuyên trách (300 DPI)", width=Inches(6.4))

    add_h2(doc, "3.1. Chức năng của Lễ tân (Receptionist)")
    rec_steps = [
        ("Bước 1: Đăng nhập hệ thống: ", "Truy cập màn hình đăng nhập, nhập tài khoản 'receptionist' và mật khẩu. Giao diện chuyển hướng vào Receptionist Dashboard."),
        ("Bước 2: Tiếp đón và Tìm kiếm bệnh nhân: ", "Nhập số CCCD hoặc SĐT của người bệnh vào ô tìm kiếm nhanh. Nếu người bệnh đã có hồ sơ, thông tin hiển thị ngay lập tức. Nếu là bệnh nhân mới, nhấn nút 'Đăng ký bệnh nhân', điền đầy đủ họ tên, ngày sinh, địa chỉ, mã thẻ BHYT và ghi nhận tiền sử dị ứng thuốc nếu có."),
        ("Bước 3: Đặt lịch khám bệnh: ", "Vào mục 'Lịch hẹn' -> Nhấn 'Tạo lịch hẹn mới'. Chọn người bệnh, chọn bác sĩ chuyên khoa và chọn khung giờ khám. Hệ thống tự động kiểm tra ca trực và báo động đỏ nếu bác sĩ hoặc phòng khám đã có lịch hẹn trùng giờ. Nếu hợp lệ, nhấn 'Lưu lịch hẹn'."),
        ("Bước 4: Quản lý hàng đợi phòng khám: ", "Khi người bệnh có mặt tại phòng khám, lễ tân chuyển trạng thái từ 'Đã hẹn' sang 'Đang chờ khám', hệ thống tự động xếp số thứ tự vào hàng đợi của bác sĩ tương ứng.")
    ]
    for title, desc in rec_steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r2 = p.add_run(desc)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)

    add_h2(doc, "3.2. Chức năng của Bác sĩ (Doctor)")
    doc_steps = [
        ("Bước 1: Tiếp nhận ca khám từ hàng đợi: ", "Bác sĩ đăng nhập tài khoản 'doctor', xem danh sách người bệnh đang chờ khám tại phòng của mình. Nhấn nút 'Tiến hành khám' trên hàng đợi."),
        ("Bước 2: Xem tóm tắt bệnh án từ Trợ lý AI (Pre-visit Briefing): ", "Nhấn nút 'Xem tóm tắt AI' trên giao diện EMR. Trợ lý AI tự động trích xuất các lần khám cũ, hiển thị thẻ tóm tắt súc tích về diễn tiến bệnh, thuốc đã dùng và dị ứng thuốc trong thời gian dưới 15 giây."),
        ("Bước 3: Ghi nhận sinh hiệu lâm sàng: ", "Đo và nhập huyết áp, mạch, nhiệt độ, SpO2, chiều cao, cân nặng. Hệ thống tự động tính chỉ số khối cơ thể (BMI) và hiển thị phân loại thể trạng."),
        ("Bước 4: Chẩn đoán bệnh theo mã ICD-10: ", "Gõ tên bệnh lý hoặc triệu chứng vào ô chẩn đoán (VD: Viêm phế quản), hệ thống gợi ý danh sách mã ICD-10 chuẩn quốc tế để bác sĩ lựa chọn."),
        ("Bước 5: Kê đơn thuốc điện tử và Cảnh báo dị ứng: ", "Tìm và thêm thuốc từ kho dược phòng khám, nhập số lượng và liều dùng. Nếu bác sĩ vô tình kê loại thuốc có hoạt chất trùng tiền sử dị ứng của người bệnh, hệ thống lập tức hiển thị cảnh báo đỏ nguy hiểm để ngăn chặn sai sót."),
        ("Bước 6: Sinh hướng dẫn sau khám và Hoàn tất ca khám: ", "Nhấn nút 'Tạo dặn dò AI', hệ thống sinh lịch uống thuốc chi tiết theo buổi (sáng, trưa, tối) và lời khuyên ăn uống. Bác sĩ kiểm tra, chỉnh sửa nếu cần và bấm 'Hoàn thành ca khám'. Hồ sơ tự động chuyển sang phân hệ kế toán.")
    ]
    for title, desc in doc_steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r2 = p.add_run(desc)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)

    add_h2(doc, "3.3. Chức năng của Kế toán (Accountant)")
    acc_steps = [
        ("Bước 1: Tiếp nhận hồ sơ chờ thanh toán: ", "Kế toán đăng nhập tài khoản 'accountant', mở danh sách hồ sơ khám đã hoàn thành cần thu tiền."),
        ("Bước 2: Kiểm tra bảng kê chi phí và BHYT: ", "Mở chi tiết phiếu khám. Hệ thống tự động tổng hợp tiền công khám và tiền thuốc trong đơn. Dựa trên mã thẻ BHYT, hệ thống tự động khấu trừ 80% hoặc 100% chi phí được bảo hiểm chi trả, hiển thị rõ ràng số tiền người bệnh phải nộp."),
        ("Bước 3: Thu tiền và Thanh toán VietQR: ", "Nếu người bệnh thanh toán chuyển khoản, kế toán bấm 'Tạo mã VietQR'. Màn hình hiển thị mã QR động chuẩn NAPAS chứa chính xác số tiền cần thu. Người bệnh dùng ứng dụng ngân hàng quét mã."),
        ("Bước 4: Xác nhận thanh toán và In biên lai: ", "Sau khi tiền vào tài khoản phòng khám, kế toán nhấn 'Xác nhận thanh toán', hóa đơn chuyển trạng thái 'Đã thanh toán'. Nhấn 'In hóa đơn' để in biên lai viện phí bàn giao cho người bệnh.")
    ]
    for title, desc in acc_steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r2 = p.add_run(desc)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)

    add_h2(doc, "3.4. Chức năng của Quản trị viên (Admin)")
    adm_steps = [
        ("Quản lý tài khoản và phân quyền: ", "Tạo tài khoản nhân sự mới, gán đúng vai trò theo ma trận RBAC, đặt lại mật khẩu và khóa tài khoản khi có biến động nhân sự."),
        ("Phân ca làm việc bác sĩ: ", "Xếp lịch trực cho từng bác sĩ theo các ca sáng, chiều, tối trong tuần và phân bổ phòng khám chuyên khoa tương ứng."),
        ("Quản lý danh mục kho dược: ", "Nhập kho thuốc mới, điều chỉnh đơn giá niêm yết, theo dõi số lượng tồn kho và thiết lập hạn mức cảnh báo sắp hết thuốc."),
        ("Báo cáo thống kê quản trị: ", "Theo dõi biểu đồ doanh thu theo ngày/tháng, số lượt người bệnh đến khám, tỷ lệ chi trả BHYT và năng suất làm việc của đội ngũ bác sĩ."),
        ("Giám sát nhật ký an ninh (Audit Logs): ", "Tra cứu lịch sử truy cập hồ sơ bệnh án nhạy cảm và nhật ký gọi API trợ lý AI để bảo đảm an toàn dữ liệu theo Nghị định 13/2023/NĐ-CP.")
    ]
    for title, desc in adm_steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r2 = p.add_run(desc)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)

    # 4. KỊCH BẢN DEMO 8 BƯỚC ĐÁNH GIÁ ĐỒ ÁN
    add_h1(doc, "4. KỊCH BẢN DEMO 8 BƯỚC ĐÁNH GIÁ ĐỒ ÁN (THỜI LƯỢNG 10 - 15 PHÚT)")
    add_p(
        doc,
        "Kịch bản trình diễn luồng nghiệp vụ khép kín từ tiếp đón đến báo cáo doanh thu phục vụ Hội đồng đánh giá chấm điểm:"
    )

    demo_steps = [
        ("Bước 1 (1 phút): Khởi chạy và Giới thiệu tổng quan", "Khởi động ứng dụng từ script 'run_app.bat'. Giới thiệu kiến trúc phân tầng FastAPI + React Vite và 4 phân hệ tương ứng 4 vai trò."),
        ("Bước 2 (2 phút): Phân hệ Lễ tân - Đăng ký bệnh nhân", "Đăng nhập tài khoản 'receptionist'. Tiếp đón người bệnh Trần Văn Bình (CCCD: 001099012345). Ghi nhận tiền sử dị ứng thuốc Penicillin."),
        ("Bước 3 (2 phút): Trình diễn thuật toán chống trùng lịch", "Đặt lịch khám vào lúc 08:30 cho Bác sĩ Lê Văn Cường. Thử đặt thêm một lịch hẹn khác cùng thời điểm 08:30 để chứng minh hệ thống phát hiện xung đột và từ chối lưu lịch."),
        ("Bước 4 (3 phút): Phân hệ Bác sĩ - Bàn khám EMR và Thẻ AI Pre-visit", "Đăng nhập tài khoản 'doctor'. Mở hồ sơ bệnh nhân từ hàng đợi. Bấm nút 'Xem tóm tắt AI' để hiển thị bản tóm tắt hồ sơ cũ không để lộ PII kèm cảnh báo y tế."),
        ("Bước 5 (2 phút): Nhập sinh hiệu và Cảnh báo dị ứng thuốc", "Nhập dấu hiệu sinh tồn (HA: 130/85, Mạch: 80, SpO2: 98%), hệ thống tự tính BMI. Nhập chẩn đoán ICD-10 'J02.9'. Thử kê thuốc Amoxicillin, hệ thống hiển thị cảnh báo đỏ dị ứng Penicillin. Đổi sang thuốc an toàn."),
        ("Bước 6 (1 phút): AI Dặn dò xuất viện", "Nhấn nút 'Tạo dặn dò sau khám', AI tự động tạo lịch uống thuốc và chế độ ăn uống. Bác sĩ hoàn thành ca khám."),
        ("Bước 7 (2 phút): Phân hệ Kế toán - Viện phí và Thanh toán VietQR", "Đăng nhập tài khoản 'accountant'. Mở hóa đơn, kiểm tra mức khấu trừ BHYT 80%. Bấm hiển thị mã thanh toán VietQR động NAPAS. Xác nhận thanh toán và in biên lai thu tiền."),
        ("Bước 8 (2 phút): Phân hệ Admin - Thống kê và Nhật ký Audit Logs", "Đăng nhập tài khoản 'admin'. Xem biểu đồ doanh thu cập nhật tức thời. Mở trang Audit Logs minh chứng toàn bộ thao tác vừa thực hiện đã được ghi vết an toàn.")
    ]

    for title, desc in demo_steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"• {title}: ")
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        r2 = p.add_run(desc)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2(doc, "4.1. Bảng tiêu chí đánh giá kịch bản Demo (Rubric chấm điểm)")
    add_p(
        doc,
        "Bảng đối soát tiêu chí đánh giá dành cho Giảng viên hướng dẫn và Hội đồng chấm điểm đồ án:"
    )

    rubric_data = [
        ("1", "Khởi động và Kiến trúc", "Khởi chạy 1-Click thành công, Swagger UI sẵn sàng, giao diện hiển thị chuẩn", "1.0 điểm", "Đạt"),
        ("2", "Tiếp đón và BHYT", "Tìm kiếm CCCD nhanh, đăng ký bệnh nhân mới, lưu số thẻ BHYT", "1.5 điểm", "Đạt"),
        ("3", "Chống trùng lịch hẹn", "ConflictChecker phát hiện xung đột thời gian thực, trả về HTTP 409", "1.5 điểm", "Đạt"),
        ("4", "Lâm sàng EMR & AI Pre-visit", "Tính BMI tự động, tra cứu ICD-10, tóm tắt bệnh sử khử PII trong < 15s", "2.0 điểm", "Đạt"),
        ("5", "Kê đơn & Cảnh báo dị ứng", "Cảnh báo đỏ tương tác dị ứng tiền sử thuốc, trừ tồn kho chính xác", "1.5 điểm", "Đạt"),
        ("6", "Viện phí & VietQR", "Khấu trừ BHYT 80%, sinh mã QR động NAPAS 247, in biên lai", "1.5 điểm", "Đạt"),
        ("7", "Quản trị & Audit Logs", "Dashboard biểu đồ doanh thu, xem nhật ký kiểm toán và an toàn AI", "1.0 điểm", "Đạt")
    ]
    tbl_rubric = doc.add_table(rows=len(rubric_data) + 1, cols=5)
    tbl_rubric.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_rubric.autofit = False
    col_w_rb = [Inches(0.6), Inches(1.8), Inches(2.7), Inches(1.0), Inches(0.6)]
    for row in tbl_rubric.rows:
        for idx, w in enumerate(col_w_rb):
            row.cells[idx].width = w

    headers_rb = ["STT", "Hạng mục đánh giá", "Tiêu chuẩn nghiệm thu chức năng", "Điểm tối đa", "Kết quả"]
    for i, h in enumerate(headers_rb):
        tbl_rubric.cell(0, i).text = h
    format_row(tbl_rubric.rows[0], HEX_HEADER_BG, RGBColor(255, 255, 255), bold=True, is_header=True)

    for r_idx, (rb_stt, rb_name, rb_std, rb_pts, rb_res) in enumerate(rubric_data):
        row = tbl_rubric.rows[r_idx + 1]
        row.cells[0].text = rb_stt
        row.cells[1].text = rb_name
        row.cells[2].text = rb_std
        row.cells[3].text = rb_pts
        row.cells[4].text = rb_res
        bg = HEX_ROW_ALT if (r_idx % 2 == 1) else "FFFFFF"
        format_row(row, bg, RGBColor(15, 23, 42), bold=False, is_header=False)
    set_table_borders(tbl_rubric)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 5. SỔ TAY CHẨN ĐOÁN VÀ KHẮC PHỤC SỰ CỐ
    add_h1(doc, "5. SỔ TAY CHẨN ĐOÁN VÀ KHẮC PHỤC SỰ CỐ (TROUBLESHOOTING)")
    add_p(
        doc,
        "Hướng dẫn nhanh cách xử lý các tình huống bất thường trong quá trình vận hành hệ thống tại phòng khám:"
    )

    issues_data = [
        ("Sự cố 1: Không kết nối được CSDL (Database Connection Refused)",
         "Nguyên nhân: Dịch vụ MySQL/PostgreSQL chưa được bật hoặc cấu hình sai biến môi trường DATABASE_URL trong tệp .env.\n"
         "Cách xử lý: Kiểm tra trạng thái dịch vụ bằng lệnh 'docker ps' hoặc 'services.msc'. Đảm bảo cổng 3306 (MySQL) hoặc 5432 (Postgres) đang mở và kiểm tra chuỗi kết nối trong backend/.env."),

        ("Sự cố 2: Lỗi gọi Trợ lý AI khi mất kết nối Internet",
         "Nguyên nhân: Dịch vụ Google Gemini API không thể truy cập do mạng Internet gián đoạn hoặc API Key hết hạn ngạch (Quota Exceeded).\n"
         "Cách xử lý: Hệ thống tự động kích hoạt cơ chế MockFallbackEngine cục bộ để tạo câu trả lời mẫu theo luật nội bộ. Người quản trị có thể thay thế GEMINI_API_KEY mới trong backend/.env mà không cần dừng hệ thống."),

        ("Sự cố 3: Đặt lịch báo trùng giờ bất thường dù bác sĩ chưa có lịch",
         "Nguyên nhân: Ca trực (DoctorShift) của bác sĩ trong ngày đó chưa được Quản trị viên kích hoạt hoặc giờ đặt nằm ngoài khung giờ trực.\n"
         "Cách xử lý: Đăng nhập quyền Admin, vào mục 'Phân ca bác sĩ', kiểm tra ca trực tương ứng đã được gán trạng thái SCHEDULED và đúng buồng khám."),

        ("Sự cố 4: Giao diện trắng trang do cổng mạng bị xung đột",
         "Nguyên nhân: Cổng 5173 hoặc 8000 đã bị chiếm dụng bởi ứng dụng khác đang chạy ngầm trên máy tính.\n"
         "Cách xử lý: Chạy tệp 'kill_ports.bat' có sẵn trong thư mục dự án để giải phóng cổng, sau đó khởi chạy lại bằng 'run_app.bat'."),

        ("Sự cố 5: Mã VietQR không hiển thị trên màn hình thu ngân",
         "Nguyên nhân: Kế toán chưa cấu hình thông tin tài khoản ngân hàng thụ hưởng của phòng khám (BANK_BIN và ACCOUNT_NO).\n"
         "Cách xử lý: Mở file backend/.env, điền chính xác mã ngân hàng thụ hưởng (VD: MB hoặc 970422) và số tài khoản chính xác để máy chủ sinh link VietQR hợp lệ.")
    ]

    for title, fix in issues_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"• {title}:\n")
        r1.font.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r2 = p.add_run(fix)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10)

    save_document(doc, "07_GenAI_SoftwareDevelopment_user-guide.docx")

if __name__ == "__main__":
    build_user_guide()
